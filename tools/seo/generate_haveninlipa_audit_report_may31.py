#!/usr/bin/env python3
"""
Generate the May 31, 2026 Full SEO Audit & Strategy Report DOCX for HavenInLipa.com
Prepared by: NetCoreSolutions
Comparison chain: Apr 7 -> Apr 15 -> May 8 -> May 31, 2026 (this report)

Scope: Fresh Full Audit & Strategy run (053126).
Strategy-continuity rule (Cedric, 2026-05-31): keep the existing 6-cluster keyword
tracker; outline ONLY existing tracker gap keywords (no net-new strategy).
This report reflects ONLY approved outputs (STOP POINT #1 and #2 both passed).

Reuses the styling/helpers established in generate_haveninlipa_audit_report_may8.py.
"""

import os
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

# Brand Colors (NetCoreSolutions)
CORAL = RGBColor(0xFF, 0x53, 0x71)
FOREST_GREEN = RGBColor(0x3B, 0x53, 0x23)
DARK_TEXT = RGBColor(0x2C, 0x2C, 0x2C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF2, 0xF2, 0xF2)
MED_GRAY = RGBColor(0x66, 0x66, 0x66)
ACCENT_RED = RGBColor(0xD9, 0x3B, 0x3B)
ACCENT_ORANGE = RGBColor(0xE8, 0x8A, 0x1A)
ACCENT_GREEN_LIGHT = RGBColor(0x0D, 0x8A, 0x4E)

BASE_DIR = "/Volumes/Files and Cloud/Dropbox/VSCode (non-Git)/seo"
LOGO_PATH = os.path.join(BASE_DIR, "Brand Assets", "NCS Logo.png")
OUTPUT_DIR = os.path.join(BASE_DIR, "content for HavenInLipa")
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "HavenInLipa_Full_SEO_Strategy_Report_May31.docx")
OUTPUT_FILE_LATEST = os.path.join(OUTPUT_DIR, "HavenInLipa_Full_SEO_Strategy_Report.docx")


# ---------- helpers (mirrored from the May 8 generator) ----------
def set_cell_shading(cell, hex_color):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_text(cell, text, bold=False, color=None, size=10, alignment=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if alignment:
        p.alignment = alignment
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if color:
        run.font.color.rgb = color
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)


def add_styled_table(doc, headers, rows, col_widths=None, header_bg="FF5371"):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, header_bg)
        set_cell_text(cell, h, bold=True, color=WHITE, size=9)
    for r_idx, row in enumerate(rows):
        bg = "F2F2F2" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            set_cell_shading(cell, bg)
            set_cell_text(cell, str(val), size=9)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
    return table


def add_heading_styled(doc, text, level=1, color=None):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "Calibri"
        if color:
            run.font.color.rgb = color
    return h


def add_body(doc, text, bold=False, italic=False, color=None, size=10.5):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if color:
        run.font.color.rgb = color
    return p


def add_bullet(doc, text, level=0, bold_prefix=""):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5 + level * 0.25)
    if bold_prefix:
        run_b = p.add_run(bold_prefix)
        run_b.bold = True
        run_b.font.size = Pt(10)
        run_b.font.name = "Calibri"
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.name = "Calibri"
    return p


def add_divider(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("-" * 80)
    run.font.size = Pt(6)
    run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)


def generate_report():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    # ---------------- COVER PAGE ----------------
    cover_para = doc.add_paragraph()
    cover_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_para.paragraph_format.space_before = Pt(80)
    if os.path.exists(LOGO_PATH):
        run = cover_para.add_run()
        run.add_picture(LOGO_PATH, width=Inches(2.5))

    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para.paragraph_format.space_before = Pt(40)
    run = title_para.add_run("FULL SEO AUDIT & STRATEGY REPORT")
    run.bold = True
    run.font.size = Pt(28)
    run.font.name = "Calibri"
    run.font.color.rgb = DARK_TEXT

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Fresh Audit Re-baseline + Continuity Strategy + Phase 3 Content")
    run.italic = True
    run.font.size = Pt(13)
    run.font.color.rgb = MED_GRAY
    run.font.name = "Calibri"

    client_para = doc.add_paragraph()
    client_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    client_para.paragraph_format.space_before = Pt(30)
    run = client_para.add_run("haveninlipa.com")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = CORAL
    run.font.name = "Calibri"

    prep = doc.add_paragraph()
    prep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prep.paragraph_format.space_before = Pt(100)
    run = prep.add_run("Prepared by: NetCoreSolutions")
    run.font.size = Pt(11)
    run.font.color.rgb = DARK_TEXT
    run.font.name = "Calibri"

    date_para = doc.add_paragraph()
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = date_para.add_run("Date: May 31, 2026")
    run.font.size = Pt(11)
    run.font.color.rgb = DARK_TEXT
    run.font.name = "Calibri"

    baseline = doc.add_paragraph()
    baseline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = baseline.add_run("Baseline Chain: Apr 7 -> Apr 15 -> May 8 -> May 31, 2026")
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = MED_GRAY
    run.font.name = "Calibri"

    conf = doc.add_paragraph()
    conf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    conf.paragraph_format.space_before = Pt(20)
    run = conf.add_run("CONFIDENTIAL")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = MED_GRAY
    run.font.name = "Calibri"

    doc.add_page_break()

    # ---------------- TABLE OF CONTENTS ----------------
    add_heading_styled(doc, "Table of Contents", level=1, color=CORAL)
    toc_items = [
        "1. Executive Summary",
        "2. Website Audit (May 31 State)",
        "    2.1 Live Site Structure & Properties",
        "    2.2 Content Quality",
        "    2.3 Technical SEO Findings",
        "    2.4 Crawl Reliability Note",
        "3. Competitor Analysis",
        "    3.1 Local Accommodation Competitors (with websites)",
        "    3.2 Who Owns the Informational SERPs",
        "    3.3 Positioning Opportunities",
        "4. Content Strategy (Continuity-First)",
        "    4.1 The Existing 6-Cluster Framework",
        "    4.2 Tracker Reconciliation Findings",
        "    4.3 Target Keyword / Topic Set (Approved)",
        "    4.4 Priority Content Plan",
        "5. SERP & Outline Strategy (Phase 2)",
        "    5.1 SERP Patterns by Gap Keyword",
        "    5.2 Superior Outlines (Summary)",
        "6. Content Production (Phase 3)",
        "    6.1 Deliverables Summary",
        "    6.2 Approach & Differentiation",
        "7. Recommendations & Next Steps",
        "    7.1 Tier A - Technical & Hygiene (Now)",
        "    7.2 Tier B - Drain Tracker Gaps",
        "    7.3 Tier C - Sustain Weekly Cadence",
        "    7.4 Tier D - Authority Backlog (Q3)",
        "    7.5 KPIs to Track",
        "8. Remediation Status (as of June 2, 2026)",
        "    8.1 Findings Close-Out Log",
        "    8.2 Client-Owned Operational Items",
        "    8.3 Net Status",
    ]
    for item in toc_items:
        p = doc.add_paragraph()
        run = p.add_run(item)
        run.font.size = Pt(10.5)
        run.font.name = "Calibri"
        p.paragraph_format.space_after = Pt(2)
    doc.add_page_break()

    # ---------------- 1. EXECUTIVE SUMMARY ----------------
    add_heading_styled(doc, "1. Executive Summary", level=1, color=CORAL)
    add_body(doc,
        "This Full SEO Audit & Strategy Report is a May 31, 2026 fresh re-baseline of HavenInLipa.com, "
        "run end-to-end (Phases 1-3) with both approval gates passed. Its defining decision is "
        "STRATEGIC CONTINUITY: at the client's direction, this run does NOT invent a new strategy. It "
        "keeps the existing 6-cluster keyword tracker as the single source of truth and limits new work "
        "to keyword gaps that were already earmarked in that tracker.")
    add_body(doc, "Headline findings:", bold=True)
    add_bullet(doc, "Both money listings are live again. The 1BR (sleeps 5, solar, P2,000) and 2BR "
                    "(sleeps 9, P2,800) were reactivated on May 31; placeholder demo listings seen earlier "
                    "in the crawl were confirmed by the client as non-final.", bold_prefix="Properties: ")
    add_bullet(doc, "Property pages are content-rich and on-brand, well above the local-competitor norm.",
               bold_prefix="Content: ")
    add_bullet(doc, "No JSON-LD schema anywhere; property pages are missing from sitemap.xml; FAQ answers "
                    "lack internal links. These are the highest-ROI fixes and need no new content.",
               bold_prefix="Technical: ")
    add_bullet(doc, "No local accommodation competitor (Lakeview Resort, JET Hotel, etc.) runs a content "
                    "engine. HIL is the only Lipa lodging brand ranking on planning intent - a durable moat.",
               bold_prefix="Moat: ")
    add_bullet(doc, "The keyword tracker's target URLs are stale (old planned slugs vs. live SEO slugs), "
                    "producing false 'Not Indexed' rows. A tracker URL-sync is queued.",
               bold_prefix="Reconciliation: ")
    add_body(doc,
        "Phase 3 produced three publish-ready drafts mapped to existing tracker gaps: a barkada / 9-pax "
        "house guide (Article #27), a Batangas road-trip hub that bases the trip in Lipa (#28), and an "
        "evergreen 'Summer in Lipa' cool-highland piece (#29). Two lower-priority gaps (senior-friendly, "
        "solo travel) were outlined and queued for Q3.")
    add_divider(doc)

    # ---------------- 2. WEBSITE AUDIT ----------------
    add_heading_styled(doc, "2. Website Audit (May 31 State)", level=1, color=CORAL)

    add_heading_styled(doc, "2.1 Live Site Structure & Properties", level=2, color=FOREST_GREEN)
    add_body(doc, "Main site nav: Home, Properties, About, Location, Blog, FAQ, Contact, Book Now. "
                  "Two live listings (reactivated 2026-05-31):")
    add_styled_table(doc,
        ["Property", "Config", "Sleeps", "Price/night", "URL"],
        [
            ["Cozy 1BR Haven (Solar, Netflix, Wi-Fi)", "1BR", "5", "P2,000", "/properties/cozy-1-bedroom"],
            ["Spacious 2BR Getaway (Wi-Fi, Parking)", "2BR", "9", "P2,800", "/properties/spacious-2-bedroom"],
        ],
        col_widths=[2.4, 0.8, 0.7, 1.0, 2.2])
    add_body(doc, "'Mickey in Lipa' is NOT yet a live listing - it exists only as the "
                  "/mickey-in-lipa-coming-soon/ blog teaser. Per client direction it remains parked until "
                  "the listing and real photos exist. Blog is healthy: 11 live posts (sitemap lastmod "
                  "2026-05-26), consistent with the project record.", italic=True)

    add_heading_styled(doc, "2.2 Content Quality", level=2, color=FOREST_GREEN)
    add_bullet(doc, "Property pages are deep (~2,800-3,000 words est.) with descriptions, amenities, "
                    "neighborhood radius guide, verified reviews, house rules, transparent PHP pricing, "
                    "per-page FAQ, and internal blog links.")
    add_bullet(doc, "Voice is on-brand: anti-hype, specific (e.g. '400 Mbps tested', 'solar backup'), "
                    "parent-to-parent / host-to-traveler.")
    add_bullet(doc, "Blog content engine is on a sustained weekly cadence.")

    add_heading_styled(doc, "2.3 Technical SEO Findings", level=2, color=FOREST_GREEN)
    add_styled_table(doc,
        ["Finding", "Severity", "Fix", "Status (as of Jun 2, 2026)"],
        [
            ["No JSON-LD schema on property pages (LodgingBusiness/Product, Review, Offer) or FAQ",
             "High", "Add structured data; no content change needed",
             "RESOLVED - schema was already live (stale-crawl false alarm); Offer (makesOffer) + additionalType added. Rich Results: 15 valid items, eligible; /faq 3 valid, 0 issues."],
            ["Property pages absent from sitemap.xml (only Home/FAQ/About/Privacy/Terms listed)",
             "High", "Add /properties/* and resubmit in GSC",
             "RESOLVED - sitemap is dynamic; both /properties/* confirmed present (stale-crawl false alarm)."],
            ["FAQ answers contain no internal links to property/blog pages",
             "Medium", "Add contextual internal links",
             "FIXED - contextual links to /#properties + property pages now live on /faq (JSON-LD text kept link-free)."],
            ["No embedded map on property pages (address as text only)",
             "Low-Med", "Embed Google Map",
             "FIXED - area-level embed added (lazy-loaded, no exact pin)."],
            ["/tag/ and /author/ archives still indexable (thin)",
             "Low", "noindex, follow",
             "FIXED - noindex,follow set in Yoast 6/2; drops on next recrawl."],
        ],
        col_widths=[2.4, 0.7, 1.7, 2.1])

    add_heading_styled(doc, "2.4 Crawl Reliability Note", level=2, color=FOREST_GREEN)
    add_body(doc,
        "Early homepage fetches returned placeholder demo listings ('The Lipa Retreat', 'Casa Verde', "
        "'The Urban Suite'). The client confirmed these were placeholders and reactivated the two real "
        "listings during the session; a cache-busted re-fetch then returned the correct state. Granular "
        "extraction figures (exact photo/review/word counts) are treated as approximate.", italic=True)
    add_divider(doc)

    # ---------------- 3. COMPETITOR ANALYSIS ----------------
    add_heading_styled(doc, "3. Competitor Analysis", level=1, color=CORAL)
    add_body(doc, "Competitor basis (client-selected): local short-term-rental hosts / resorts. On "
                  "investigation the landscape splits into two tiers.")

    add_heading_styled(doc, "3.1 Local Accommodation Competitors (with websites)", level=2, color=FOREST_GREEN)
    add_styled_table(doc,
        ["Competitor", "Type", "Own blog / guides?", "Direct booking?", "Depth"],
        [
            ["Lakeview Resort", "Villa resort, infinity pools", "None", "Contact form", "Thin-Mod"],
            ["JET Hotel", "3-star city hotel", "None", "Strong", "Thin-Mod"],
            ["The Farm at San Benito", "Luxury wellness resort", "Limited (PR)", "Own engine", "Brand-led"],
            ["Cintai Corito's Garden", "Balinese resort hotel", "Minimal", "Own engine", "Thin"],
        ],
        col_widths=[1.8, 2.0, 1.4, 1.2, 0.9])
    add_body(doc, "Key insight: NOT ONE local accommodation competitor runs a content/SEO engine targeting "
                  "'things to do in Lipa', area guides, or trip planning. This is HavenInLipa's structural "
                  "moat - it is the only Lipa lodging brand that ranks (and can keep ranking) on "
                  "informational and planning intent.", bold=True)

    add_heading_styled(doc, "3.2 Who Owns the Informational SERPs", level=2, color=FOREST_GREEN)
    add_bullet(doc, "Pure STR hosts (Joenice, Gillera, PMPJ, etc.) exist only as OTA listings - no website, "
                    "no content.")
    add_bullet(doc, "Informational SERPs are dominated by OTAs/aggregators (Tripadvisor, Expedia, Trip.com, "
                    "Traveloka, Vrbo, Yelp) and PH travel blogs (Tara Lets Anywhere, Pinoy Adventurista, "
                    "Trip101). Generic, non-local, no booking stake in Lipa.")

    add_heading_styled(doc, "3.3 Positioning Opportunities", level=2, color=FOREST_GREEN)
    add_bullet(doc, "'Where to stay in Lipa' decision content - OTAs only list, nobody helps travelers choose.")
    add_bullet(doc, "Group/barkada and long-stay intent - resorts chase events/golf; nobody owns these from "
                    "a host's POV. Maps directly to the live 2BR (groups) and 1BR (long-stay/solar).")
    add_divider(doc)

    # ---------------- 4. CONTENT STRATEGY ----------------
    add_heading_styled(doc, "4. Content Strategy (Continuity-First)", level=1, color=CORAL)

    add_heading_styled(doc, "4.1 The Existing 6-Cluster Framework", level=2, color=FOREST_GREEN)
    add_body(doc, "Per client direction, the strategy keeps the six clusters already defined in "
                  "HavenInLipa_Keyword_Tracker.xlsx. No new cluster framework was created.")
    add_styled_table(doc,
        ["#", "Cluster", "Dominant Intent"],
        [
            ["1", "Short-Term Rentals", "Transactional"],
            ["2", "Things to Do in Lipa", "Informational"],
            ["3", "Weekend Getaway / Audience", "Mixed / Commercial / Seasonal"],
            ["4", "Travel Planning & Logistics", "Informational"],
            ["5", "Direct Booking & Value", "Commercial"],
            ["6", "Branded / Property (parked - Mickey)", "Branded"],
        ],
        col_widths=[0.5, 3.5, 2.8])

    add_heading_styled(doc, "4.2 Tracker Reconciliation Findings", level=2, color=FOREST_GREEN)
    add_bullet(doc, "Stale target URLs: the tracker's Page Indexing Tracker checks old planned slugs "
                    "(e.g. /things-to-do-lipa-city/) that 404; live slugs are longer/SEO-optimized "
                    "(e.g. /15-best-things-to-do-in-lipa-city-batangas-2026-locals-guide/). Several "
                    "'Not Indexed' rows are false alarms. Sync queued. [RESOLVED 6/2 - all 11 live "
                    "blog slugs synced to their live URLs and cross-checked against GSC's 22 indexed "
                    "pages; false 'Not Indexed' rows cleared.]", bold_prefix="Finding 1: ")
    add_bullet(doc, "The tracker predates the June-August calendar (last reviewed 2026-05-18) - it has no "
                    "rows for articles #15-26. Add them during the URL-sync. [RESOLVED 6/2 - all 12 "
                    "articles #15-26 added to Keyword Master with 3 keywords each (36 rows) plus 12 "
                    "Page-Indexing pipeline rows; clusters, target slugs, and weekly publish dates set.]",
               bold_prefix="Finding 2: ")

    add_heading_styled(doc, "4.3 Target Keyword / Topic Set (Approved)", level=2, color=FOREST_GREEN)
    add_body(doc, "Approved scope (STOP POINT #1): existing tracker GAP keywords only. Net-new keywords "
                  "were declined in favor of strict continuity.")
    add_styled_table(doc,
        ["Keyword", "Cluster", "Priority", "Tracker Status"],
        [
            ["barkada group trip lipa", "3", "Medium", "Gap -> Article #27 (drafted)"],
            ["batangas road trip itinerary", "4", "Medium", "Gap -> Article #28 (drafted)"],
            ["summer in lipa (evergreen)", "3", "Medium", "Gap -> Article #29 (drafted)"],
            ["senior-friendly lipa getaway", "3", "Low", "Gap - outlined, Q3 backlog"],
            ["solo travel guide to lipa", "3", "Low", "Gap - outlined, Q3 backlog"],
        ],
        col_widths=[2.6, 0.8, 1.0, 2.8])

    add_heading_styled(doc, "4.4 Priority Content Plan", level=2, color=FOREST_GREEN)
    add_styled_table(doc,
        ["Priority", "Item", "Tie", "Effort", "Impact", "Status (Jun 2)"],
        [
            ["P0", "Property + FAQ schema (JSON-LD)", "gates Cluster 1", "Low", "High", "DONE 6/2"],
            ["P0", "Add /properties/* to sitemap + resubmit GSC", "gates Cluster 1", "Low", "High", "DONE 6/2"],
            ["P0", "Sync tracker URLs + add #15-26 rows", "tracker upkeep", "Low", "High", "DONE 6/2"],
            ["P1", "Barkada / 9-pax guide (#27)", "C3 existing gap", "Med", "High", "Drafted; publish ~Sep 7"],
            ["P2", "Batangas road-trip hub (#28)", "C4 existing gap", "Med", "Med", "Drafted; publish ~Sep 14"],
            ["P3", "Summer in Lipa evergreen (#29)", "C3 existing gap", "Low", "Med", "Drafted; publish ~Sep 21"],
            ["Backlog", "Senior-friendly / solo travel", "C3 Tier-D gaps", "Med", "Low", "Outlined (Q3)"],
        ],
        col_widths=[0.7, 2.3, 1.3, 0.5, 0.5, 1.6])
    add_divider(doc)

    # ---------------- 5. SERP & OUTLINE STRATEGY ----------------
    add_heading_styled(doc, "5. SERP & Outline Strategy (Phase 2)", level=1, color=CORAL)

    add_heading_styled(doc, "5.1 SERP Patterns by Gap Keyword", level=2, color=FOREST_GREEN)
    add_styled_table(doc,
        ["Keyword", "SERP reality", "The wedge (gap exploited)"],
        [
            ["barkada group trip lipa", "Beach-skewed listicles (Anilao/Laiya)",
             "Cost-per-head math (P2,800/9 ~ P311); cool inland alt; drives 2BR"],
            ["batangas road trip itinerary", "Lipa treated as a pass-through",
             "'Use Lipa as your base' hub; ties #1/#4/#5/#7/#10"],
            ["summer in lipa", "Resort lists; ignore elevation",
             "22-34C highland = anti-heat/anti-crowd; solar 1BR angle"],
            ["senior-friendly lipa getaway", "No Lipa-specific senior content",
             "Flat heritage walks, mild climate, hospitals, calm"],
            ["solo travel guide to lipa", "Thin/scattered (TikTok, small blogs)",
             "Safe, walkable, coffee-culture; crosslinks Work From Lipa"],
        ],
        col_widths=[1.9, 2.4, 3.0])

    add_heading_styled(doc, "5.2 Superior Outlines (Summary)", level=2, color=FOREST_GREEN)
    add_body(doc, "Five superior outlines were built (full H1-H3 detail in 053126/Phase2_SERP_and_Outlines.md). "
                  "Each meets the internal-linking floor: >=2 blog links, >=1 property page, the Book Direct "
                  "article (#6), and the standardized footer - using verified live slugs.")
    add_bullet(doc, "Two carry-forward flags: (1) 'summer in lipa' built evergreen (no year) since the "
                    "2026 summer window has closed - dated promo push deferred to Feb-Mar 2027; (2) the "
                    "senior piece will not claim the 1BR is senior/PWD-friendly until its layout is "
                    "fact-checked.")
    add_divider(doc)

    # ---------------- 6. CONTENT PRODUCTION ----------------
    add_heading_styled(doc, "6. Content Production (Phase 3)", level=1, color=CORAL)

    add_heading_styled(doc, "6.1 Deliverables Summary", level=2, color=FOREST_GREEN)
    add_styled_table(doc,
        ["#", "Title", "Cluster", "Status"],
        [
            ["27", "Barkada Getaway in Lipa - Whole House vs Beach Resort", "3", "Drafted (publish-ready)"],
            ["28", "Batangas Road Trip Itinerary - Lipa as Base", "4", "Drafted (publish-ready)"],
            ["29", "Summer in Lipa - Cool Highland Escape (evergreen)", "3", "Drafted (publish-ready)"],
            ["-", "Senior-Friendly Lipa Getaway", "3", "Outline only (Q3)"],
            ["-", "Solo Travel Guide to Lipa", "3", "Outline only (Q3)"],
        ],
        col_widths=[0.5, 3.7, 0.8, 2.0])
    add_body(doc, "All three drafts follow the single-file convention (blogger guide embedded under the "
                  "'STOP - DO NOT COPY' separator) and the featured-image filename convention. Files live "
                  "in content for HavenInLipa/053126/.", italic=True)

    add_heading_styled(doc, "6.2 Approach & Differentiation", level=2, color=FOREST_GREEN)
    add_bullet(doc, "Barkada (#27): leads with honest cost-per-head math (~P311/head) - the number every "
                    "competitor listicle omits. Drives the 2BR sleeps-9.")
    add_bullet(doc, "Road trip (#28): repositions Lipa from pass-through to base; the densest internal-link "
                    "hub, wiring the existing travel articles together.")
    add_bullet(doc, "Summer (#29): leans the solar / brownout-proof 1BR differentiator - a real PH-summer "
                    "advantage no competitor frames.")
    add_bullet(doc, "All retain the locked brand voice: PHP budgets, named local businesses, anti-hype, "
                    "direct-booking conversion lever.")
    add_divider(doc)

    # ---------------- 7. RECOMMENDATIONS ----------------
    add_heading_styled(doc, "7. Recommendations & Next Steps", level=1, color=CORAL)

    add_heading_styled(doc, "7.1 Tier A - Technical & Hygiene (Now)", level=2, color=FOREST_GREEN)
    add_bullet(doc, "Add LodgingBusiness/Product + Review + Offer JSON-LD to both property pages; FAQPage "
                    "schema on /faq.")
    add_bullet(doc, "Add /properties/* to sitemap.xml and resubmit in GSC.")
    add_bullet(doc, "Sync the keyword tracker's target URLs to live slugs; add rows for #15-26.")
    add_bullet(doc, "Add internal links inside FAQ answers; embed maps on property pages.")
    add_body(doc, "Status: all Tier A technical & hygiene items above were COMPLETED on June 2, 2026. "
                  "See Section 8 for the close-out log and proof.", bold=True, color=ACCENT_GREEN_LIGHT)

    add_heading_styled(doc, "7.2 Tier B - Drain Tracker Gaps", level=2, color=FOREST_GREEN)
    add_bullet(doc, "Publish #27 (barkada), #28 (road-trip hub), #29 (summer evergreen) - all drafted.")
    add_bullet(doc, "Wire the hub: add in-body links back to #28 from #1/#4/#5/#7/#10.")

    add_heading_styled(doc, "7.3 Tier C - Sustain Weekly Cadence", level=2, color=FOREST_GREEN)
    add_bullet(doc, "Resume the locked weekly-Monday cadence with the already-drafted #15-26; slot the new "
                    "gap articles after the run or interleave where a slot opens.")
    add_bullet(doc, "Run the R1 'Things to Do' quarterly refresh as pre-scheduled (Aug 31).")

    add_heading_styled(doc, "7.4 Tier D - Authority Backlog (Q3)", level=2, color=FOREST_GREEN)
    add_bullet(doc, "Draft the senior-friendly and solo-travel pieces (outlined). Fact-check 1BR "
                    "accessibility before publishing the senior article.")

    add_heading_styled(doc, "7.5 KPIs to Track", level=2, color=FOREST_GREEN)
    add_styled_table(doc,
        ["KPI", "Source", "Cadence"],
        [
            ["Impressions / clicks per tracked keyword", "GSC Performance", "Weekly (Mon)"],
            ["Avg. position movement on HIGH-priority terms", "GSC", "Weekly"],
            ["Indexing status (live slugs, not stale)", "GSC URL Inspection", "Per publish"],
            ["Rich-result eligibility after schema ship", "GSC Enhancements", "Monthly"],
            ["Direct-booking inquiries vs. OTA", "Client / contact form", "Monthly"],
        ],
        col_widths=[3.2, 2.0, 1.5])
    add_divider(doc)

    # ---------------- 8. REMEDIATION STATUS (ADDENDUM) ----------------
    add_heading_styled(doc, "8. Remediation Status (as of June 2, 2026)", level=1, color=CORAL)
    add_body(doc,
        "This section is a post-audit close-out addendum, added June 2, 2026. The findings and "
        "recommendations above are preserved as the May 31 point-in-time snapshot; this log records "
        "their resolution. Source: client remediation report 'SEO_Resolution_2026-06-02.md'.")
    add_body(doc,
        "Important context: the two P0 'High' technical findings (missing schema, property pages "
        "absent from the sitemap) were, on verification, FALSE ALARMS from a stale crawl that hit "
        "placeholder demo listings ('The Lipa Retreat', 'Casa Verde', 'The Urban Suite') - see the "
        "Crawl Reliability Note in 2.4. Schema was already live and the sitemap was already dynamic; "
        "both were independently re-validated. The genuinely actionable items (FAQ links, map, "
        "archive noindex, base-URL canonicalization, and the tracker reconciliation) were shipped.",
        italic=True)

    add_heading_styled(doc, "8.1 Findings Close-Out Log", level=2, color=FOREST_GREEN)
    add_styled_table(doc,
        ["Audit Finding", "Severity", "Status", "Proof / Notes"],
        [
            ["JSON-LD schema on property pages + FAQ (2.3)", "High", "RESOLVED",
             "Already live (false alarm); added Offer (makesOffer) + additionalType. Google Rich Results: 15 valid items, eligible; /faq 3 valid, 0 issues. Commits 918c5e5, a65ef16, 4b300b1."],
            ["Property pages in sitemap.xml (2.3)", "High", "RESOLVED",
             "Sitemap is dynamic; both /properties/* URLs confirmed present. Stale-crawl false alarm."],
            ["FAQ answers need internal links (2.3)", "Medium", "FIXED",
             "Contextual links to /#properties + property pages now live on /faq; JSON-LD text kept link-free."],
            ["Map on property pages (2.3)", "Low-Med", "FIXED",
             "Area-level Google Map embed added; lazy-loaded, no exact pin."],
            ["/tag/ + /author/ archives indexable (2.3)", "Low", "FIXED",
             "noindex,follow set in Yoast 6/2; will drop from index on next recrawl."],
            ["Base-URL www vs non-www mismatch", "-", "FIXED",
             "Fallback corrected to non-www; 301 www -> apex (and .vercel.app -> apex) verified 6/2."],
            ["Tracker stale target URLs (4.2 Finding 1)", "High (hygiene)", "RESOLVED",
             "All 11 live blog slugs synced to live URLs; cross-checked vs GSC's 22 indexed pages."],
            ["Tracker missing #15-26 rows (4.2 Finding 2)", "High (hygiene)", "RESOLVED",
             "36 keyword rows (#15-26, 3 each) + 12 indexing rows added; clusters, slugs, weekly dates set."],
        ],
        col_widths=[2.2, 0.9, 0.9, 2.9])

    add_heading_styled(doc, "8.2 Client-Owned Operational Items", level=2, color=FOREST_GREEN)
    add_bullet(doc, "Indexing requested for the 3 live-but-unindexed posts (#11 family-staycation, "
                    "#12 romantic-getaway, #13 Mickey teaser) - submitted in GSC 6/2; verify 'Indexed' "
                    "within 1-7 days.")
    add_bullet(doc, "Article #14 (Family Weekend Without Beach Crowds) published Mon 6/1.")
    add_bullet(doc, "301 www -> apex redirect verified 6/2 (Vercel).")
    add_bullet(doc, "#15-26 topics/keywords supplied and loaded into the tracker (12 articles x 3 "
                    "keywords) - 6/2.")
    add_bullet(doc, "GSC reconciliation confirmed #10 (How to Get to Lipa) is LIVE and indexed - the "
                    "earlier '404' concern is resolved.")

    add_heading_styled(doc, "8.3 Net Status", level=2, color=FOREST_GREEN)
    add_body(doc,
        "No audit items remain open. Every code-actionable finding is shipped and validated, and the "
        "keyword tracker is synced to live + GSC reality. Remaining activity is execution on the "
        "existing plan - not audit remediation: publish #15-26 on the weekly Monday cadence (Jun 8 -> "
        "Aug 24), run the R1 refresh (Aug 31), and publish the new gap articles #27/#28/#29 in "
        "September - requesting indexing per publish and syncing each live slug back into the tracker.",
        bold=True)
    add_divider(doc)

    closing = doc.add_paragraph()
    closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = closing.add_run("Prepared by NetCoreSolutions  |  May 31, 2026  (Remediation addendum: Jun 2, 2026)  |  CONFIDENTIAL")
    run.font.size = Pt(9)
    run.font.color.rgb = MED_GRAY
    run.font.name = "Calibri"

    doc.save(OUTPUT_FILE)
    # Refresh the unsuffixed "latest" copy per folder convention
    doc.save(OUTPUT_FILE_LATEST)
    print("Saved:", OUTPUT_FILE)
    print("Saved (latest):", OUTPUT_FILE_LATEST)


if __name__ == "__main__":
    generate_report()
