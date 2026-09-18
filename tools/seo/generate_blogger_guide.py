#!/usr/bin/env python3
"""
Generate a professional DOCX Blogger Guide for HavenInLipa — Articles 3-10.
"""

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

# ── Colors ────────────────────────────────────────────────────────────────────

CORAL = RGBColor(0xFF, 0x53, 0x71)
FOREST_GREEN = RGBColor(0x3B, 0x53, 0x23)
DARK_TEXT = RGBColor(0x33, 0x33, 0x33)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY_HEX = "F2F2F2"
CORAL_HEX = "FF5371"
FOREST_HEX = "3B5323"


# ── Helpers (adapted from generate_report.py) ────────────────────────────────

def set_cell_shading(cell, hex_color):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_text(cell, text, bold=False, color=None, size=10, alignment=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if alignment:
        p.alignment = alignment
    run = p.add_run(str(text))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if color:
        run.font.color.rgb = color
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)


def add_styled_table(doc, headers, rows, header_color=CORAL_HEX, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, header_color)
        set_cell_text(cell, h, bold=True, color=WHITE, size=9)

    for r_idx, row in enumerate(rows):
        bg = LIGHT_GRAY_HEX if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            set_cell_shading(cell, bg)
            set_cell_text(cell, str(val), size=9)

    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)

    return table


def add_heading_styled(doc, text, level=1, color=None):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "Calibri"
        if color:
            run.font.color.rgb = color
    return h


def add_body(doc, text, bold=False, italic=False, color=None, size=10.5):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if color:
        run.font.color.rgb = color
    return p


def add_bullet(doc, text, level=0, bold_prefix=""):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5 + level * 0.25)
    if bold_prefix:
        run_b = p.add_run(bold_prefix)
        run_b.bold = True
        run_b.font.size = Pt(10)
        run_b.font.name = "Calibri"
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.name = "Calibri"
    return p


def add_divider(doc, color_hex=CORAL_HEX):
    """Add a colored horizontal divider line."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("━" * 80)
    run.font.size = Pt(7)
    r, g, b = int(color_hex[0:2], 16), int(color_hex[2:4], 16), int(color_hex[4:6], 16)
    run.font.color.rgb = RGBColor(r, g, b)


def add_metadata_table(doc, metadata_dict):
    """Add a 2-column metadata table (Label | Value)."""
    rows_data = list(metadata_dict.items())
    table = doc.add_table(rows=len(rows_data), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    for i, (label, value) in enumerate(rows_data):
        bg = LIGHT_GRAY_HEX if i % 2 == 0 else "FFFFFF"
        cell_label = table.rows[i].cells[0]
        cell_value = table.rows[i].cells[1]
        set_cell_shading(cell_label, FOREST_HEX)
        set_cell_text(cell_label, label, bold=True, color=WHITE, size=9.5)
        set_cell_shading(cell_value, bg)
        set_cell_text(cell_value, value, size=9.5)
        cell_label.width = Inches(1.5)
        cell_value.width = Inches(5.0)

    return table


# ── Article Data ──────────────────────────────────────────────────────────────

ARTICLES = [
    {
        "num": 3,
        "name": "Weekend Getaway in Lipa City, Batangas — Your Chill Escape Near Manila",
        "slug": "weekend-getaway-lipa-city-batangas",
        "category": "Travel & Itineraries",
        "tags": "Lipa City, Batangas, weekend getaway, near Manila, itinerary, staycation",
        "featured_image": 'A scenic road or highway view heading south (SLEX/STAR Tollway) with green Batangas hills, or a cozy cafe scene in Lipa City',
        "featured_alt": "Weekend getaway in Lipa City Batangas - chill escape near Manila",
        "excerpt": "Planning a weekend getaway near Manila? Lipa City, Batangas is just 1 hour away — no crowds, great food, and mountain views. See our 2-day itinerary and budget breakdown!",
        "images": [
            ("1", 'SLEX or STAR Tollway road with green hills in the background — the "road trip" vibe', "1200 x 630px", "Road trip to Lipa City via SLEX and STAR Tollway - 1 hour from Manila", "After the intro paragraph (before \"Why Lipa Is Manila's Best-Kept Weekend Secret\")"),
            ("2", "A beautifully plated dish at a restaurant or a cozy restaurant interior", "1200 x 800px", "Lunch at Casa Marikit restaurant in Lipa City - number one on TripAdvisor", "Inside \"Day 1\" (after the Casa Marikit lunch entry)"),
            ("3", "A cup of barako coffee with beans on the side, or the Cafe de Lipa storefront", "1200 x 800px", "Barako coffee at Cafe de Lipa - afternoon coffee stop on a weekend getaway", "Inside \"Day 1\" (after the Cafe de Lipa entry)"),
            ("4", "Hikers at the Mt. Maculot Rockies viewpoint overlooking Taal Lake", "1200 x 800px", "Mt Maculot Rockies viewpoint - day 2 adventure option on a Lipa weekend getaway", "Inside \"Day 2: Adventure Mode\" (after the Mt. Maculot option)"),
            ("5", "Interior or exterior photo of a HavenInLipa property", "1200 x 800px", "HavenInLipa vacation rental - best place to stay for a Lipa City weekend getaway", "Inside \"Where to Stay in Lipa City\" section"),
            ("6", "A simple infographic or flat-lay photo of travel essentials (wallet, phone, keys, coffee)", "1200 x 630px", "Weekend getaway budget breakdown - Lipa City for 2 people costs under 9000 pesos", "Inside \"Budget Breakdown\" section (above or below the table)"),
        ],
    },
    {
        "num": 4,
        "name": "Mt. Maculot Hiking Guide 2026 — Trail Tips, Routes + Where to Stay in Lipa",
        "slug": "mt-maculot-hiking-guide",
        "category": "Outdoor Adventures",
        "tags": "Mt Maculot, hiking, Batangas, Cuenca, day hike, trail guide, Taal Lake",
        "featured_image": "The iconic Rockies viewpoint at Mt. Maculot with a hiker sitting on the edge, Taal Lake in the background",
        "featured_alt": "Mt Maculot Rockies viewpoint overlooking Taal Lake - hiking guide 2026",
        "excerpt": "Complete Mt. Maculot hiking guide for 2026. Three trail routes explained, step-by-step itinerary, budget breakdown, what to bring, and the best base camp in nearby Lipa City.",
        "images": [
            ("1", "Wide panoramic shot of Taal Lake from the Rockies — the hero landscape shot", "1200 x 630px", "Panoramic view of Taal Lake from Mt Maculot Rockies in Batangas", "After the intro paragraph (before \"Quick Stats\")"),
            ("2", "The rocky ridge trail with hikers, showing the exposed rock face", "1200 x 800px", "Mt Maculot Rockies Trail - easiest route with the best views", "Inside \"Route 1: Rockies Trail\" section"),
            ("3", "Hiker on a forested trail section or the rope-assisted section", "1200 x 800px", "Mt Maculot full traverse trail - rope section on the way to the summit", "Inside \"Route 3: Full Traverse\" section"),
            ("4", "Sunrise or morning light over Taal Lake from the Rockies viewpoint", "1200 x 800px", "Sunrise at Mt Maculot Rockies viewpoint - morning light on Taal Lake", "Inside \"Step-by-Step Itinerary\" (at the \"7:30 AM — Reach the Rockies\" entry)"),
            ("5", "Flat-lay of hiking gear: shoes, water bottle, trail food, sunscreen, hat, backpack", "1200 x 800px", "Mt Maculot hiking essentials - what to bring checklist", "Inside \"What to Bring\" section (above the checklist)"),
            ("6", "HavenInLipa property photo — interior or exterior", "1200 x 800px", "HavenInLipa in Lipa City - base camp 15 minutes from Mt Maculot jump-off", "Inside \"Where to Stay Near Mt. Maculot\" section"),
        ],
    },
    {
        "num": 5,
        "name": "Best Restaurants & Cafes in Lipa City, Batangas — 2026 Food Guide",
        "slug": "best-restaurants-lipa-city-batangas",
        "category": "Lipa City Guide",
        "tags": "Lipa City, restaurants, where to eat, Batangas, lomi, barako coffee, food guide",
        "featured_image": "A close-up of a bowl of lomi noodles or a spread of Filipino dishes on a table",
        "featured_alt": "Best restaurants and food in Lipa City Batangas - 2026 food guide",
        "excerpt": "Craving great food in Lipa City? From fine dining at Casa Marikit to legendary lomi noodles at Beegee's and barako coffee at Cafe de Lipa — here's where to eat in Lipa, Batangas.",
        "images": [
            ("1", "An appetizing spread of different Filipino dishes — or a vibrant food market scene", "1200 x 630px", "Lipa City food scene - restaurants cafes and street food in Batangas", "After the intro (before \"Lipa's Food Scene\")"),
            ("2", "Plated food at Casa Marikit (arugula salad, pizza, or steak) or the restaurant interior/garden", "1200 x 800px", "Casa Marikit restaurant Lipa City - number one rated on TripAdvisor", "Inside \"Casa Marikit\" section"),
            ("3", "Close-up of a steaming bowl of special lomi — thick noodles, egg, meat, rich broth", "1200 x 800px", "Beegees Special Lomi in Lipa City - famous Batangas lomi noodle soup", "Inside \"Beegee's Lomi House\" section"),
            ("4", "A cup of hot barako coffee with beans, or the cafe interior/storefront", "1200 x 800px", "Cafe de Lipa barako coffee - the original Batangas coffee destination", "Inside \"Cafe de Lipa\" section"),
            ("5", "Colorful street food stall or a display of kakanin, kwek-kwek, and local snacks", "1200 x 800px", "Lipa City street food and pasalubong - public market finds", "Inside \"Street Food & Pasalubong\" section"),
            ("6", "Bags of barako coffee beans, tablea, and other Batangas pasalubong items", "1200 x 800px", "Lipa City pasalubong - barako coffee beans and tablea to bring home", "Inside \"Pasalubong Must-Buys\" section"),
        ],
    },
    {
        "num": 6,
        "name": "Why Book Direct Instead of Airbnb? (A Philippine Host's Honest Take)",
        "slug": "why-book-direct-instead-of-airbnb",
        "category": "Booking Tips",
        "tags": "book direct, Airbnb, save money, Philippines, vacation rental, GCash, tips",
        "featured_image": 'A split-screen concept: one side showing Airbnb fees/checkout screen, the other showing a direct message with a host — or a simple graphic showing "Book Direct = Save 20%"',
        "featured_alt": "Why book direct instead of Airbnb - save on fees with HavenInLipa",
        "excerpt": "Save P500+ per night by booking direct instead of Airbnb. Honest fee breakdown from a Philippine host, safety tips for direct bookings, and how to book HavenInLipa in Lipa City.",
        "images": [
            ("1", 'A simple graphic or infographic showing "Airbnb Fees: 14-20% Service Fee" with a peso sign', "1200 x 630px", "Airbnb service fees breakdown - 14 to 20 percent added to every booking", 'Inside "How Much Do Airbnb Fees Really Cost You?" (above the first comparison table)'),
            ("2", "Screenshot or photo of a friendly chat conversation with host Melody (or a mockup showing direct communication)", "1200 x 800px", "Book direct with HavenInLipa - talk to host Melody directly not a call center", 'Inside "7 Reasons to Book Direct" (after reason #1 or #2)'),
            ("3", "GCash and BPI logos, or a phone showing a GCash QR code payment screen", "1200 x 630px", "HavenInLipa accepts GCash Maya and BPI bank transfer - flexible Philippine payment", 'Inside "Flexible Payment Options" (reason #4)'),
            ("4", "Screenshot of HavenInLipa's Facebook page or Google listing showing real reviews and ratings", "1200 x 800px", "HavenInLipa verified reviews on Facebook and Google - 51 five star reviews", 'Inside "Is It Safe to Book Direct?" section'),
            ("5", "Screenshot of the HavenInLipa.com website homepage or booking contact form", "1200 x 800px", "How to book direct at HavenInLipa.com - step by step process", 'Inside "How to Book Direct with HavenInLipa" section'),
        ],
    },
    {
        "num": 7,
        "name": "Taal Volcano Day Trip from Lipa City — 2026 Updated Guide",
        "slug": "taal-volcano-day-trip-from-lipa",
        "category": "Travel & Itineraries",
        "tags": "Taal Volcano, day trip, Lipa City, Batangas, Taal Lake, heritage town, boat tour",
        "featured_image": "The iconic Taal Volcano island in the middle of Taal Lake — shot from Tagaytay Ridge or from a boat on the lake",
        "featured_alt": "Taal Volcano day trip from Lipa City Batangas - 2026 updated guide",
        "excerpt": "Plan your Taal Volcano day trip from Lipa City. Boat tours on the lake, Taal Heritage Town walks, scenic viewpoints, and a full budget breakdown. Updated for 2026 — know what's open.",
        "images": [
            ("1", "Wide landscape shot of Taal Volcano island sitting in Taal Lake — the classic postcard view", "1200 x 630px", "Taal Volcano in Taal Lake Batangas - view from Tagaytay Ridge", 'After the intro (before "Taal Volcano in 2026")'),
            ("2", "A bangka (outrigger boat) on Taal Lake with Volcano Island in the background", "1200 x 800px", "Taal Lake boat tour from Talisay - bangka with Volcano Island view", 'Inside "Taal Lake Boat Tour" section'),
            ("3", "The Basilica of Saint Martin de Tours — the massive white church facade", "1200 x 800px", "Basilica of Saint Martin de Tours in Taal Heritage Town - largest church in Philippines", 'Inside "Taal Heritage Town Walking Tour" section'),
            ("4", "A heritage street in Taal with ancestral houses — Spanish-Filipino architecture", "1200 x 800px", "Taal Heritage Town ancestral houses - centuries old Filipino Spanish homes", 'Inside "Taal Heritage Town" (after the ancestral houses mention)'),
            ("5", "The Taal view from a Tagaytay viewpoint or restaurant — volcano, lake, and sky", "1200 x 800px", "Taal Volcano view from Tagaytay Ridge - scenic viewpoint on a day trip from Lipa", 'Inside "Scenic Viewpoints from Tagaytay Ridge" section'),
            ("6", "HavenInLipa property interior — living room or bedroom, cozy and inviting", "1200 x 800px", "HavenInLipa vacation rental - your Lipa City base for Taal Volcano day trips", 'Inside "Combine It: Taal Day Trip + Lipa Staycation" section'),
        ],
    },
    {
        "num": 8,
        "name": "Holy Week in Lipa City, Batangas — A Peaceful Retreat Near Manila",
        "slug": "holy-week-getaway-lipa-city-batangas",
        "category": "Travel & Itineraries",
        "tags": "Holy Week, Semana Santa, Lipa City, Batangas, Visita Iglesia, church, retreat",
        "featured_image": "San Sebastian Cathedral in Lipa City during golden hour or with candle-lit procession — or a serene church interior",
        "featured_alt": "Holy Week in Lipa City Batangas - peaceful retreat near Manila",
        "excerpt": "Spend Holy Week in Lipa City, Batangas. Historic churches, Visita Iglesia, quiet reflection, cool highland air, and zero beach crowds — just 1 hour from Manila.",
        "images": [
            ("1", "San Sebastian Cathedral exterior — grand facade, ideally at golden hour or dusk", "1200 x 630px", "San Sebastian Cathedral Lipa City - heart of Holy Week celebrations in Batangas", 'After the intro (before "Why Lipa is Perfect for Semana Santa")'),
            ("2", "The Carmelite Monastery entrance or chapel interior — peaceful, serene atmosphere", "1200 x 800px", "Carmelite Monastery Lipa City - site of the Marian apparitions and Shower of Roses", 'Inside "Deep Catholic Heritage" (after the Carmelite Monastery mention)'),
            ("3", "Interior of a historic Batangas church with candles or religious imagery — solemn atmosphere", "1200 x 800px", "Visita Iglesia in Lipa City - church hopping during Holy Week in Batangas", 'Inside "Visita Iglesia" section'),
            ("4", "Sunrise view from Mt. Maculot or hikers on the trail at dawn", "1200 x 800px", "Early morning hike at Mt Maculot during Holy Week - sunrise over Taal Lake", 'Inside "Early Morning Mt. Maculot Hike" section'),
            ("5", "HavenInLipa property photo — cozy, restful interior (bedroom or living area)", "1200 x 800px", "HavenInLipa rental in Lipa City - peaceful accommodation for Holy Week retreat", 'Inside "Where to Stay: HavenInLipa" section'),
        ],
    },
    {
        "num": 9,
        "name": "Work From Lipa: The Affordable Remote Work Staycation Near Manila",
        "slug": "remote-work-staycation-lipa-city",
        "category": "Work & Extended Stays",
        "tags": "remote work, staycation, WFH, digital nomad, Lipa City, fast WiFi, work from anywhere",
        "featured_image": 'A laptop on a desk near a window with a green/nature view outside — the "work from anywhere" aesthetic',
        "featured_alt": "Remote work staycation in Lipa City - work from anywhere with 400 Mbps WiFi",
        "excerpt": "Ditch Manila rent. Work remotely from Lipa City with 400 Mbps WiFi, a full kitchen, and Netflix — starting at P2,000 per night. Quiet, affordable, and just 1 hour from the metro.",
        "images": [
            ("1", "A clean workspace setup: laptop, coffee cup, notebook, on a desk with natural light", "1200 x 630px", "Remote work setup at HavenInLipa - 400 Mbps WiFi workspace in Lipa City", 'After the intro (before "Why Remote Workers Love Lipa City")'),
            ("2", "Screenshot of a WiFi speed test showing 400 Mbps (or close to it)", "800 x 600px", "WiFi speed test at HavenInLipa Cozy 1BR Haven - 400 Mbps fiber connection", 'Inside "Internet you can rely on" paragraph in the WFH Setup section'),
            ("3", "The kitchen area of a HavenInLipa unit — stove, cookware, clean counters", "1200 x 800px", "Full kitchen at HavenInLipa - cook your own meals during extended stays", 'Inside "A full kitchen so you can cook your own meals" paragraph'),
            ("4", "Sunset or golden hour view from a mountain or hilltop in Batangas", "1200 x 800px", "Sunset hike at Mt Maculot after work - remote work life in Lipa City", 'Inside "What to Do After Work Hours" (after the Mt. Maculot entry)'),
            ("5", "HavenInLipa property exterior or living room — welcoming, homey feel", "1200 x 800px", "HavenInLipa extended stay rental - weekly and monthly rates available in Lipa City", 'Inside "Extended Stay Options and Pricing" section'),
        ],
    },
    {
        "num": 10,
        "name": "How to Get to Lipa City from Manila — Complete 2026 Guide",
        "slug": "how-to-get-to-lipa-from-manila",
        "category": "Getting Here",
        "tags": "Lipa City, Manila, bus, SLEX, STAR Tollway, commute, directions, transport",
        "featured_image": 'A highway/expressway shot (SLEX or STAR Tollway) with a green "Lipa" exit sign, or a bus at a terminal',
        "featured_alt": "How to get to Lipa City from Manila - drive bus or Grab 2026 guide",
        "excerpt": "Getting to Lipa Batangas from Manila is easy. Drive via SLEX in 1 hour, take a bus for P238, or book a Grab. Full 2026 routes, fares, and pro tips inside.",
        "images": [
            ("1", "A simple map graphic showing Manila to SLEX to STAR Tollway to Lipa City route with distance/time labels", "1200 x 630px", "Map showing route from Manila to Lipa City via SLEX and STAR Tollway", 'After the intro (before "Quick Summary" table)'),
            ("2", "SLEX or STAR Tollway road — open highway, clear day, heading south", "1200 x 800px", "Driving to Lipa City via SLEX and STAR Tollway - about 1 hour from Manila", 'Inside "By Car" section'),
            ("3", "PITX terminal building exterior or bus boarding area", "1200 x 800px", "PITX Paranaque Integrated Terminal - take ALPS bus to Lipa City for 238 pesos", 'Inside "From PITX" section under "By Bus"'),
            ("4", "A Lipa City tricycle (the local transport) on a street", "1200 x 800px", "Lipa City tricycle - local transport from terminal to your accommodation", 'Inside "From Lipa Terminal to Your Accommodation" section'),
            ("5", "HavenInLipa property exterior — welcoming entrance, well-lit", "1200 x 800px", "HavenInLipa vacation rental in Lipa City - your home base in Batangas", 'Inside "Ready to Visit Lipa City?" section at the end'),
        ],
    },
]

# Quick reference summary rows
QUICK_REF = [
    ("3", "Weekend Getaway in Lipa City", "weekend-getaway-lipa-city-batangas", "Travel & Itineraries", "weekend getaway, Lipa City, near Manila"),
    ("4", "Mt. Maculot Hiking Guide 2026", "mt-maculot-hiking-guide", "Outdoor Adventures", "Mt Maculot, hiking, Batangas"),
    ("5", "Best Restaurants in Lipa City", "best-restaurants-lipa-city-batangas", "Lipa City Guide", "restaurants, Lipa City, lomi"),
    ("6", "Book Direct vs Airbnb", "why-book-direct-instead-of-airbnb", "Booking Tips", "book direct, Airbnb, save money"),
    ("7", "Taal Volcano Day Trip", "taal-volcano-day-trip-from-lipa", "Travel & Itineraries", "Taal Volcano, day trip, boat tour"),
    ("8", "Holy Week Getaway", "holy-week-getaway-lipa-city-batangas", "Travel & Itineraries", "Holy Week, Semana Santa, retreat"),
    ("9", "Remote Work Staycation", "remote-work-staycation-lipa-city", "Work & Extended Stays", "remote work, staycation, fast WiFi"),
    ("10", "How to Get to Lipa", "how-to-get-to-lipa-from-manila", "Getting Here", "Manila, bus, directions"),
]

IMAGE_SOURCES = [
    ("Your own photos", "HavenInLipa property, kitchen, workspace, food", "Free"),
    ("Unsplash / Pexels", "Mountains, landscapes, highways, sunsets", "Free"),
    ("Canva", "Map graphics, infographics, fee breakdowns", "Free tier"),
    ("Google Maps screenshot", "Route maps", "Free"),
    ("Facebook pages", "Restaurant food photos, cafe interiors (ask permission)", "Free"),
    ("AI image generators", "Generic lifestyle shots, flat-lays", "Varies"),
]

WP_STEPS = [
    "Click where you want the image in your content",
    "Type /image and press Enter",
    "Click Upload and select your file",
    "After uploading, click the image -> go to Block tab in right sidebar -> fill in Alt Text",
    "Keep image file names lowercase with hyphens (e.g., cafe-de-lipa-barako.jpg)",
]


# ── Main Generator ────────────────────────────────────────────────────────────

def generate():
    doc = Document()

    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # Default font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    style.font.color.rgb = DARK_TEXT

    # ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
    # COVER PAGE
    # ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
    for _ in range(6):
        doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("HavenInLipa")
    run.bold = True
    run.font.size = Pt(36)
    run.font.color.rgb = CORAL
    run.font.name = "Calibri"

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Blog Posting Guide")
    run.bold = True
    run.font.size = Pt(28)
    run.font.color.rgb = FOREST_GREEN
    run.font.name = "Calibri"

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("All Articles (3-10)")
    run.font.size = Pt(18)
    run.font.color.rgb = DARK_TEXT
    run.font.name = "Calibri"

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Instructions for Blogger")
    run.italic = True
    run.font.size = Pt(14)
    run.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    run.font.name = "Calibri"

    doc.add_paragraph()

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("April 7, 2026")
    run.font.size = Pt(12)
    run.font.color.rgb = DARK_TEXT
    run.font.name = "Calibri"

    doc.add_page_break()

    # ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
    # GENERAL INSTRUCTIONS
    # ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
    add_heading_styled(doc, "How to Use This Guide", level=1, color=CORAL)
    add_body(doc, "Use this guide when posting each article to WordPress. For each article, set the metadata in the Post sidebar, add the tags, select the category, upload the featured image, and place images in the correct sections.")
    add_body(doc, 'Reminder: In the article files, ## = H2 heading, ### = H3 heading. The title (H1) is already set in the WordPress Title field. Never add another H1 inside the content.', bold=True)

    add_divider(doc)

    # ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
    # QUICK REFERENCE TABLE
    # ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
    add_heading_styled(doc, "Quick Reference: All Articles", level=1, color=CORAL)

    add_styled_table(
        doc,
        headers=["#", "Article Name", "Slug", "Category", "Tags (Top 3)"],
        rows=QUICK_REF,
        header_color=CORAL_HEX,
        col_widths=[0.4, 2.0, 2.0, 1.3, 1.6],
    )

    doc.add_page_break()

    # ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
    # EACH ARTICLE
    # ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
    for art in ARTICLES:
        # Coral divider at top
        add_divider(doc, CORAL_HEX)

        # Article heading
        add_heading_styled(doc, f"ARTICLE {art['num']}: {art['name']}", level=1, color=CORAL)

        # Metadata table
        add_heading_styled(doc, "Post Metadata", level=2, color=FOREST_GREEN)
        add_metadata_table(doc, {
            "URL Slug": art["slug"],
            "Category": art["category"],
            "Tags": art["tags"],
            "Excerpt": art["excerpt"],
        })

        doc.add_paragraph()

        # Featured image
        add_heading_styled(doc, "Featured Image", level=2, color=FOREST_GREEN)
        add_body(doc, f"Description: {art['featured_image']}", bold=True)
        add_body(doc, f'Alt Text: "{art["featured_alt"]}"', italic=True)

        doc.add_paragraph()

        # Image placement guide table
        add_heading_styled(doc, "Image Placement Guide", level=2, color=FOREST_GREEN)

        img_rows = []
        for img in art["images"]:
            img_rows.append((img[0], img[1], img[2], img[3], img[4]))

        add_styled_table(
            doc,
            headers=["#", "What to Photograph", "Size", "Alt Text", "Where to Place"],
            rows=img_rows,
            header_color=CORAL_HEX,
            col_widths=[0.3, 1.8, 0.8, 1.8, 2.0],
        )

        doc.add_page_break()

    # ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
    # APPENDIX
    # ━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━
    add_divider(doc, FOREST_HEX)
    add_heading_styled(doc, "Appendix", level=1, color=CORAL)

    # Where to Get Images
    add_heading_styled(doc, "Where to Get Images", level=2, color=FOREST_GREEN)
    add_styled_table(
        doc,
        headers=["Source", "Best For", "Cost"],
        rows=IMAGE_SOURCES,
        header_color=FOREST_HEX,
        col_widths=[1.8, 3.5, 1.0],
    )

    doc.add_paragraph()

    # How to Add Images in WordPress
    add_heading_styled(doc, "How to Add Images in WordPress", level=2, color=FOREST_GREEN)
    for i, step in enumerate(WP_STEPS, 1):
        add_body(doc, f"{i}. {step}")

    # ── Save ──
    output_path = "/Users/cedricpcastillo/Documents/VSCode/seo/content for HavenInLipa/HavenInLipa_Blogger_Guide.docx"
    doc.save(output_path)
    print(f"Generated: {output_path}")


if __name__ == "__main__":
    generate()
