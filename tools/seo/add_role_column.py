#!/usr/bin/env python3
"""
Add a "Role" (Primary / Secondary) column to Keyword Master, inserted right
after "Priority", and populate every data row.

Concept: each target PAGE gets exactly ONE primary keyphrase (the head term that
takes the Yoast focus keyphrase / title / H1 / slug / first-100-words), and the
other keywords mapping to that same page are secondaries. Role is a per-keyword
attribute distinct from Priority (which is strategic importance across the whole
tracker), so it sits next to Priority.

Safe in-place patcher:
  - Backs up to HavenInLipa_SEO_Tracker.pre-role.backup.xlsx first.
  - Keyword Master has no formulas / data validations / conditional formatting /
    tables / merges (verified), so insert_cols is clean here.
  - Inserting a column before the metric columns shifts the columns that the
    'Cluster KPIs' sheet formulas reference (Position/Impressions/Clicks), so this
    script REBUILDS those formulas afterwards using the new column positions.
  - Column widths are re-applied by header name (robust to openpyxl version
    differences in whether insert_cols moves column_dimensions).

Run only when the workbook is CLOSED in Excel.
"""

import os
import shutil
import sys
from copy import copy

from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

BASE_DIR = "/Volumes/Files and Cloud/Dropbox/VSCode (non-Git)/seo"
CLIENT_FOLDER = os.path.join(BASE_DIR, "content for HavenInLipa")
FILE = os.path.join(CLIENT_FOLDER, "HavenInLipa_SEO_Tracker.xlsx")
BACKUP = os.path.join(CLIENT_FOLDER, "HavenInLipa_SEO_Tracker.pre-role.backup.xlsx")
KM_SHEET = "Keyword Master"
CK_SHEET = "Cluster KPIs"
ROLE_WIDTH = 12

# --------------------------------------------------------------------------- #
# Role assignment: rows that are the PRIMARY keyphrase for their page.
# Everything else in the data range is Secondary. Derived per target-URL:
# one primary per page = the head term (best slug/title match, then priority+intent).
# --------------------------------------------------------------------------- #
PRIMARY_ROWS = {
    2,   # things to do in lipa city            -> Art#1 (things-to-do)
    5,   # weekend getaway near manila          -> Art#3
    6,   # mt maculot hiking guide              -> Art#4
    7,   # best restaurants in lipa             -> Art#5
    11,  # book direct vs airbnb                -> Art#6 (why-book-direct)  [r9 strongest 2ndary]
    14,  # taal volcano day trip from lipa      -> Art#7
    15,  # holy week getaway lipa               -> Art#8
    16,  # remote work staycation philippines   -> Art#9 (work-from-lipa)
    17,  # how to get to lipa from manila       -> Art#10
    18,  # family staycation lipa               -> Art#11
    22,  # romantic getaway batangas            -> Art#12
    26,  # full house rental lipa               -> Mickey property page (full/sleeps-15)
    27,  # family house rental lipa city        -> Mickey property page (family/sleeps-11)
    28,  # mickey in lipa                       -> Mickey launch/hub post (coming-soon)
    32,  # family weekend batangas              -> Art#14
    36,  # independence day long weekend lipa   -> Art#15
    39,  # what to do in lipa when it rains     -> Art#16
    42,  # lipa barako coffee                   -> Art#17
    45,  # best lomi in lipa                    -> Art#18
    48,  # work from lipa                       -> Art#19 (rainy-season)
    51,  # lipa vs tagaytay                     -> Art#20
    54,  # casa de segunda lipa                 -> Art#21
    57,  # carmel of lipa                       -> Art#22
    60,  # august long weekend lipa 2026        -> Art#23
    63,  # lipa charter day                     -> Art#24
    66,  # ninoy aquino day long weekend 2026   -> Art#25
    69,  # national heroes day weekend lipa2026 -> Art#26
    72,  # short term rental lipa city          -> Homepage
    76,  # barkada group trip lipa              -> future Art#27 (planned page)
    77,  # summer in lipa 2026                  -> future Seasonal LP (planned page)
    78,  # senior-friendly lipa getaway         -> future senior article (planned page)
    79,  # solo travel guide to lipa            -> future solo article (planned page)
    81,  # batangas road trip itinerary         -> future hub Art#28 (planned page)
}


def main():
    if not os.path.exists(FILE):
        sys.exit(f"ERROR: workbook not found: {FILE}")
    lock = os.path.join(os.path.dirname(FILE), "~$" + os.path.basename(FILE))
    if os.path.exists(lock):
        sys.exit("ERROR: workbook appears open in Excel (~$ lock). Close it and re-run.")

    shutil.copy2(FILE, BACKUP)
    print(f"Backup created: {BACKUP}")

    wb = load_workbook(FILE)
    ws = wb[KM_SHEET]

    headers = [c.value for c in ws[1]]
    if "Role" in headers:
        sys.exit("ERROR: 'Role' column already exists. Aborting to avoid duplicate.")
    if "Priority" not in headers:
        sys.exit("ERROR: 'Priority' column not found.")
    priority_idx = headers.index("Priority") + 1        # 1-indexed
    insert_at = priority_idx + 1                          # right after Priority
    last_row = ws.max_row

    # capture widths by header name (robust re-apply after insert)
    name_width = {}
    for i in range(1, ws.max_column + 1):
        nm = ws.cell(row=1, column=i).value
        if nm is not None:
            name_width[str(nm)] = ws.column_dimensions[get_column_letter(i)].width

    # snapshot style sources BEFORE shifting
    hdr_style_src = ws.cell(row=1, column=priority_idx)   # Priority header (coral)
    body_fill_by_row = {}                                 # zebra fill per data row
    for r in range(2, last_row + 1):
        # Priority cell carries a colored pill; use Cluster cell for the row's zebra
        body_fill_by_row[r] = copy(ws.cell(row=r, column=headers.index("Cluster") + 1).fill)
    body_border = copy(ws.cell(row=2, column=headers.index("Cluster") + 1).border)

    # --- insert the column ---
    ws.insert_cols(insert_at, 1)
    col_letter = get_column_letter(insert_at)

    # header
    h = ws.cell(row=1, column=insert_at, value="Role")
    h.font = copy(hdr_style_src.font)
    h.fill = copy(hdr_style_src.fill)
    h.alignment = copy(hdr_style_src.alignment)
    h.border = copy(hdr_style_src.border)

    # body values + styling (preserve zebra)
    n_primary = 0
    for r in range(2, last_row + 1):
        val = "Primary" if r in PRIMARY_ROWS else "Secondary"
        if val == "Primary":
            n_primary += 1
        c = ws.cell(row=r, column=insert_at, value=val)
        c.font = Font(name="Calibri", bold=(val == "Primary"), size=10, color="2C2C2C")
        c.fill = body_fill_by_row[r]
        c.border = body_border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # re-apply column widths by header name; Role gets its own width
    for i in range(1, ws.max_column + 1):
        nm = ws.cell(row=1, column=i).value
        if nm == "Role":
            ws.column_dimensions[get_column_letter(i)].width = ROLE_WIDTH
        elif nm is not None and str(nm) in name_width and name_width[str(nm)] is not None:
            ws.column_dimensions[get_column_letter(i)].width = name_width[str(nm)]

    # extend auto-filter across the new column
    if ws.auto_filter.ref:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{last_row}"

    # --- repair Cluster KPIs formulas (metric columns shifted right by 1) ---
    new_headers = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}
    clu = get_column_letter(new_headers["Cluster"])
    pos = get_column_letter(new_headers["Current Avg. Position"])
    imp = get_column_letter(new_headers["Impressions (28d)"])
    clk = get_column_letter(new_headers["Clicks (28d)"])
    km = f"'{KM_SHEET}'!"
    clu_rng = f"{km}${clu}$2:${clu}${last_row}"
    pos_rng = f"{km}${pos}$2:${pos}${last_row}"
    imp_rng = f"{km}${imp}$2:${imp}${last_row}"
    clk_rng = f"{km}${clk}$2:${clk}${last_row}"

    ck = wb[CK_SHEET]
    fixed = 0
    for r in range(4, ck.max_row + 1):
        a = ck.cell(row=r, column=1).value
        if a is None:
            continue
        if str(a).startswith("TOTAL"):
            ck.cell(row=r, column=4).value = (
                f'=IFERROR(AVERAGEIFS({pos_rng},{pos_rng},">0"),"-")')
            fixed += 1
            continue
        ck.cell(row=r, column=2).value = f"=SUMIF({clu_rng},A{r},{clk_rng})"
        ck.cell(row=r, column=3).value = f"=SUMIF({clu_rng},A{r},{imp_rng})"
        ck.cell(row=r, column=4).value = (
            f'=IFERROR(AVERAGEIFS({pos_rng},{clu_rng},A{r},{pos_rng},">0"),"-")')
        ck.cell(row=r, column=5).value = f"=COUNTIF({clu_rng},A{r})"
        ck.cell(row=r, column=6).value = (
            f'=COUNTIFS({clu_rng},A{r},{pos_rng},">0",{pos_rng},"<=10")')
        fixed += 1

    wb.save(FILE)
    print(f"'Role' inserted at column {col_letter} (after Priority). "
          f"{n_primary} Primary / {last_row - 1 - n_primary} Secondary.")
    print(f"Cluster KPIs formulas repaired: {fixed} rows now reference "
          f"cluster={clu} clicks={clk} impr={imp} pos={pos}.")
    print(f"Saved: {FILE}")


if __name__ == "__main__":
    main()
