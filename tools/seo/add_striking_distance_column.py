#!/usr/bin/env python3
"""
Add a formula-driven "Striking Distance?" column to Keyword Master, inserted
right after "Current Avg. Position", so a keyword's position and its flag sit
side by side.

Striking distance = GSC average position 11-20 (page 2). These are the highest-ROI
optimization targets: the page already ranks, a refresh + internal links + a
sharper title often pushes it onto page 1. The column auto-flags them the moment
you paste this month's positions - no manual bookkeeping.

  Formula per row:  =IF(AND(<pos>>=11,<pos><=20),"Striking distance","")
  (blank/0/other positions -> blank; conditional formatting highlights the hits.)

Safe in-place patcher:
  - Backs up to HavenInLipa_SEO_Tracker.pre-striking.backup.xlsx first.
  - Inserting after Current Avg. Position shifts Impressions/Clicks, so this script
    REBUILDS the Cluster KPIs formulas afterwards using the new positions.
  - Widths re-applied by header name; zebra preserved.

Run only when the workbook is CLOSED in Excel.
"""

import os
import shutil
import sys
from copy import copy

from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import FormulaRule

BASE_DIR = "/Volumes/Files and Cloud/Dropbox/VSCode (non-Git)/seo"
CLIENT_FOLDER = os.path.join(BASE_DIR, "content for HavenInLipa")
FILE = os.path.join(CLIENT_FOLDER, "HavenInLipa_SEO_Tracker.xlsx")
BACKUP = os.path.join(CLIENT_FOLDER, "HavenInLipa_SEO_Tracker.pre-striking.backup.xlsx")
KM_SHEET = "Keyword Master"
CK_SHEET = "Cluster KPIs"

COL_NAME = "Striking Distance?"
COL_WIDTH = 18
ANCHOR = "Current Avg. Position"      # insert immediately after this column
FLAG_TEXT = "Striking distance"
HL_FILL = "FCE4CC"                    # light amber highlight
HL_TEXT = "9C5700"                    # amber text


def main():
    if not os.path.exists(FILE):
        sys.exit(f"ERROR: workbook not found: {FILE}")
    lock = os.path.join(os.path.dirname(FILE), "~$" + os.path.basename(FILE))
    if os.path.exists(lock):
        sys.exit("ERROR: workbook appears open in Excel (~$ lock). Close it and re-run.")

    shutil.copy2(FILE, BACKUP)
    print(f"Backup created: {BACKUP}")

    wb = load_workbook(FILE)
    ws = wb[KM_SHEET]
    headers = [c.value for c in ws[1]]
    if COL_NAME in headers:
        sys.exit(f"ERROR: '{COL_NAME}' already exists. Aborting.")
    if ANCHOR not in headers:
        sys.exit(f"ERROR: anchor column '{ANCHOR}' not found.")

    anchor_idx = headers.index(ANCHOR) + 1          # 1-indexed
    insert_at = anchor_idx + 1                        # right after position
    pos_letter = get_column_letter(anchor_idx)        # unchanged (left of insert)
    last_row = ws.max_row

    # widths by header name (robust re-apply)
    name_width = {}
    for i in range(1, ws.max_column + 1):
        nm = ws.cell(row=1, column=i).value
        if nm is not None:
            name_width[str(nm)] = ws.column_dimensions[get_column_letter(i)].width

    # style sources
    hdr_src = ws.cell(row=1, column=anchor_idx)
    cluster_idx = headers.index("Cluster") + 1
    body_fill_by_row = {r: copy(ws.cell(row=r, column=cluster_idx).fill)
                        for r in range(2, last_row + 1)}
    body_border = copy(ws.cell(row=2, column=cluster_idx).border)

    ws.insert_cols(insert_at, 1)
    new_letter = get_column_letter(insert_at)

    # header
    h = ws.cell(row=1, column=insert_at, value=COL_NAME)
    h.font = copy(hdr_src.font)
    h.fill = copy(hdr_src.fill)
    h.alignment = copy(hdr_src.alignment)
    h.border = copy(hdr_src.border)

    # body formulas
    for r in range(2, last_row + 1):
        c = ws.cell(row=r, column=insert_at,
                    value=f'=IF(AND({pos_letter}{r}>=11,{pos_letter}{r}<=20),"{FLAG_TEXT}","")')
        c.font = Font(name="Calibri", size=10, color="2C2C2C")
        c.fill = body_fill_by_row[r]
        c.border = body_border
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    # highlight hits (amber) via conditional formatting on the flag column
    rng = f"{new_letter}2:{new_letter}{last_row}"
    ws.conditional_formatting.add(
        rng,
        FormulaRule(formula=[f'${new_letter}2="{FLAG_TEXT}"'],
                    fill=PatternFill("solid", fgColor=HL_FILL),
                    font=Font(name="Calibri", size=10, bold=True, color=HL_TEXT)))

    # widths
    for i in range(1, ws.max_column + 1):
        nm = ws.cell(row=1, column=i).value
        if nm == COL_NAME:
            ws.column_dimensions[get_column_letter(i)].width = COL_WIDTH
        elif nm is not None and str(nm) in name_width and name_width[str(nm)] is not None:
            ws.column_dimensions[get_column_letter(i)].width = name_width[str(nm)]

    # extend auto-filter
    if ws.auto_filter.ref:
        ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{last_row}"

    # repair Cluster KPIs (Impressions/Clicks shifted right by 1; Position unchanged)
    nh = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None}
    clu = get_column_letter(nh["Cluster"])
    pos = get_column_letter(nh["Current Avg. Position"])
    imp = get_column_letter(nh["Impressions (28d)"])
    clk = get_column_letter(nh["Clicks (28d)"])
    km = f"'{KM_SHEET}'!"
    clu_rng = f"{km}${clu}$2:${clu}${last_row}"
    pos_rng = f"{km}${pos}$2:${pos}${last_row}"
    imp_rng = f"{km}${imp}$2:${imp}${last_row}"
    clk_rng = f"{km}${clk}$2:${clk}${last_row}"

    ck = wb[CK_SHEET]
    fixed = 0
    for r in range(4, ck.max_row + 1):
        a = ck.cell(row=r, column=1).value
        if a is None:
            continue
        if str(a).startswith("TOTAL"):
            ck.cell(row=r, column=4).value = f'=IFERROR(AVERAGEIFS({pos_rng},{pos_rng},">0"),"-")'
            fixed += 1
            continue
        ck.cell(row=r, column=2).value = f"=SUMIF({clu_rng},A{r},{clk_rng})"
        ck.cell(row=r, column=3).value = f"=SUMIF({clu_rng},A{r},{imp_rng})"
        ck.cell(row=r, column=4).value = f'=IFERROR(AVERAGEIFS({pos_rng},{clu_rng},A{r},{pos_rng},">0"),"-")'
        ck.cell(row=r, column=5).value = f"=COUNTIF({clu_rng},A{r})"
        ck.cell(row=r, column=6).value = f'=COUNTIFS({clu_rng},A{r},{pos_rng},">0",{pos_rng},"<=10")'
        fixed += 1

    wb.save(FILE)
    print(f"'{COL_NAME}' inserted at column {new_letter} (after {ANCHOR}); "
          f"formula flags position 11-20 across {last_row - 1} rows.")
    print(f"Cluster KPIs formulas repaired: {fixed} rows -> "
          f"cluster={clu} clicks={clk} impr={imp} pos={pos}.")
    print(f"Saved: {FILE}")


if __name__ == "__main__":
    main()
