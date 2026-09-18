#!/usr/bin/env python3
"""
Generate the July 4, 2026 SEO DELTA AUDIT REPORT DOCX for HavenInLipa.com
Prepared by: NetCoreSolutions
Baseline chain: Apr 7 -> Apr 15 -> May 8 -> May 31 -> Jun 18 -> Jul 4, 2026 (this report)

Scope: DELTA audit (folder 070426). Re-crawl of the live site checking ONLY
changes since the 2026-06-18 delta run.
Competitor re-analysis intentionally skipped per scope decision (moat unchanged,
baseline < 90 days).
Strategy-continuity rule (Cedric, 2026-05-31, STANDING): keyword tracker read
first; 6-cluster framework kept; no net-new strategy proposed.

Reuses the styling/helpers from generate_haveninlipa_audit_report_jun18.py
(cloned per convention -- prior generators are never edited in place).
"""

import os
from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

# Brand Colors (NetCoreSolutions)
CORAL = RGBColor(0xFF, 0x53, 0x71)
FOREST_GREEN = RGBColor(0x3B, 0x53, 0x23)
DARK_TEXT = RGBColor(0x2C, 0x2C, 0x2C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_GRAY = RGBColor(0xF2, 0xF2, 0xF2)
MED_GRAY = RGBColor(0x66, 0x66, 0x66)
ACCENT_RED = RGBColor(0xD9, 0x3B, 0x3B)
ACCENT_ORANGE = RGBColor(0xE8, 0x8A, 0x1A)
ACCENT_GREEN_LIGHT = RGBColor(0x0D, 0x8A, 0x4E)

BASE_DIR = "/Volumes/Files and Cloud/Dropbox/VSCode (non-Git)/seo"
LOGO_PATH = os.path.join(BASE_DIR, "Brand Assets", "NCS Logo.png")
OUTPUT_DIR = os.path.join(BASE_DIR, "content for HavenInLipa")
OUTPUT_FILE = os.path.join(OUTPUT_DIR, "HavenInLipa_SEO_Audit_Report_Jul4.docx")
OUTPUT_FILE_LATEST = os.path.join(OUTPUT_DIR, "HavenInLipa_SEO_Audit_Report.docx")
OUTPUT_FILE_RUN = os.path.join(OUTPUT_DIR, "070426", "HavenInLipa_SEO_Audit_Report_Jul4.docx")


# ---------- helpers (mirrored from the Jun 18 generator) ----------
def set_cell_shading(cell, hex_color):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def set_cell_text(cell, text, bold=False, color=None, size=10, alignment=None):
    cell.text = ""
    p = cell.paragraphs[0]
    if alignment:
        p.alignment = alignment
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if color:
        run.font.color.rgb = color
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)


def add_styled_table(doc, headers, rows, col_widths=None, header_bg="FF5371"):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, header_bg)
        set_cell_text(cell, h, bold=True, color=WHITE, size=9)
    for r_idx, row in enumerate(rows):
        bg = "F2F2F2" if r_idx % 2 == 0 else "FFFFFF"
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            set_cell_shading(cell, bg)
            set_cell_text(cell, str(val), size=9)
    if col_widths:
        for row in table.rows:
            for i, w in enumerate(col_widths):
                row.cells[i].width = Inches(w)
    return table


def add_heading_styled(doc, text, level=1, color=None):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.name = "Calibri"
        if color:
            run.font.color.rgb = color
    return h


def add_body(doc, text, bold=False, italic=False, color=None, size=10.5):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    if color:
        run.font.color.rgb = color
    return p


def add_bullet(doc, text, level=0, bold_prefix=""):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Inches(0.5 + level * 0.25)
    if bold_prefix:
        run_b = p.add_run(bold_prefix)
        run_b.bold = True
        run_b.font.size = Pt(10)
        run_b.font.name = "Calibri"
    run = p.add_run(text)
    run.font.size = Pt(10)
    run.font.name = "Calibri"
    return p


def add_divider(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run("-" * 80)
    run.font.size = Pt(6)
    run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)


def generate_report():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    # ---------------- COVER PAGE ----------------
    cover_para = doc.add_paragraph()
    cover_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_para.paragraph_format.space_before = Pt(80)
    if os.path.exists(LOGO_PATH):
        run = cover_para.add_run()
        run.add_picture(LOGO_PATH, width=Inches(2.5))

    title_para = doc.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_para.paragraph_format.space_before = Pt(40)
    run = title_para.add_run("SEO DELTA AUDIT REPORT")
    run.bold = True
    run.font.size = Pt(28)
    run.font.name = "Calibri"
    run.font.color.rgb = DARK_TEXT

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("Live-Site Re-Crawl & Remediation - Changes Since the June 18 Delta Run")
    run.italic = True
    run.font.size = Pt(13)
    run.font.color.rgb = MED_GRAY
    run.font.name = "Calibri"

    client_para = doc.add_paragraph()
    client_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    client_para.paragraph_format.space_before = Pt(30)
    run = client_para.add_run("haveninlipa.com")
    run.bold = True
    run.font.size = Pt(20)
    run.font.color.rgb = CORAL
    run.font.name = "Calibri"

    prep = doc.add_paragraph()
    prep.alignment = WD_ALIGN_PARAGRAPH.CENTER
    prep.paragraph_format.space_before = Pt(100)
    run = prep.add_run("Prepared by: NetCoreSolutions")
    run.font.size = Pt(11)
    run.font.color.rgb = DARK_TEXT
    run.font.name = "Calibri"

    date_para = doc.add_paragraph()
    date_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = date_para.add_run("Date: July 4, 2026")
    run.font.size = Pt(11)
    run.font.color.rgb = DARK_TEXT
    run.font.name = "Calibri"

    baseline = doc.add_paragraph()
    baseline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = baseline.add_run("Baseline Chain: Apr 7 -> Apr 15 -> May 8 -> May 31 -> Jun 18 -> Jul 4, 2026")
    run.italic = True
    run.font.size = Pt(10)
    run.font.color.rgb = MED_GRAY
    run.font.name = "Calibri"

    scope = doc.add_paragraph()
    scope.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = scope.add_run("Scope: Delta re-crawl (folder 070426). Changes since Jun 18. Competitor re-analysis skipped - moat unchanged.")
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = MED_GRAY
    run.font.name = "Calibri"

    conf = doc.add_paragraph()
    conf.alignment = WD_ALIGN_PARAGRAPH.CENTER
    conf.paragraph_format.space_before = Pt(20)
    run = conf.add_run("CONFIDENTIAL")
    run.bold = True
    run.font.size = Pt(9)
    run.font.color.rgb = MED_GRAY
    run.font.name = "Calibri"

    doc.add_page_break()

    # ---------------- TABLE OF CONTENTS ----------------
    add_heading_styled(doc, "Table of Contents", level=1, color=CORAL)
    toc_items = [
        "1. Executive Summary",
        "2. Website Audit (Delta - What Changed Since June 18)",
        "    2.1 New Content - Article #18 Best Lomi LIVE",
        "    2.2 Publishing Cadence - On Schedule",
        "    2.3 Main-Site Redeploy (Jul 2) - Schema Re-Verified",
        "    2.4 Archive noindex - Now Confirmed Live",
        "    2.5 NEW & RESOLVED - /book Endpoints Canonicalized to Homepage",
        "3. Keyword Tracker Reconciliation",
        "    3.1 Article #18 - Slug Sync & Status Flip",
        "    3.2 Archive Rows Reconciled",
        "    3.3 /book Rows Updated",
        "    3.4 Continuity Note",
        "4. Competitor Position (Carry-Forward)",
        "5. Findings & Resolution Log",
        "6. Recommendations & Next Steps",
        "    6.1 Immediate (Client-Owned)",
        "    6.2 Continue the Plan",
        "    6.3 KPIs to Track",
    ]
    for item in toc_items:
        p = doc.add_paragraph()
        run = p.add_run(item)
        run.font.size = Pt(10.5)
        run.font.name = "Calibri"
        p.paragraph_format.space_after = Pt(2)
    doc.add_page_break()

    # ---------------- 1. EXECUTIVE SUMMARY ----------------
    add_heading_styled(doc, "1. Executive Summary", level=1, color=CORAL)
    add_body(doc,
        "This is a DELTA audit of HavenInLipa.com, run July 4, 2026 (folder 070426). It re-crawls the live "
        "site and reports ONLY what changed since the June 18 delta run. Per the standing strategy-continuity "
        "rule, the SEO Tracker was read first; the 6-cluster framework is unchanged and no net-new strategy is "
        "proposed. Competitor re-analysis was intentionally skipped - the structural moat (no local "
        "accommodation rival runs a content engine) is unchanged and the baseline is well under 90 days old.")
    add_body(doc, "This was a quiet, healthy delta. Highlights:", bold=True)
    add_bullet(doc, "Article #18 (Best Lomi in Lipa) went live Mon Jun 29 as scheduled - HTTP 200, index/follow, "
                    "Article schema, ~2,600 words, meeting the internal-linking floor. A slug drift was caught "
                    "and corrected in the tracker (planned '-local-guide' -> live 'best-lomi-lipa-city').",
               bold_prefix="New content on schedule: ")
    add_bullet(doc, "The main site was redeployed Jul 2 (homepage/faq/about/privacy/terms re-stamped). A redeploy "
                    "is exactly when JSON-LD silently breaks - all property + FAQ schema was re-verified present. "
                    "No regressions.", bold_prefix="Redeploy clean: ")
    add_bullet(doc, "The tag/author archive noindex (set ~Jun 2) is now CONFIRMED live in raw HTML on all four "
                    "archives - they will drop from the index on next crawl.", bold_prefix="Fix confirmed live: ")
    add_bullet(doc, "The three Mickey /book form endpoints were canonicalizing to the homepage (sharing its "
                    "title and canonical) and two were indexed. Flagged and FIXED same-day: noindex,follow + "
                    "self-referential canonical + unique titles now live on all three; GSC recrawl requested.",
               bold_prefix="Found & fixed same-day: ")
    add_divider(doc)

    # ---------------- 2. WEBSITE AUDIT (DELTA) ----------------
    add_heading_styled(doc, "2. Website Audit (Delta - What Changed Since June 18)", level=1, color=CORAL)

    add_heading_styled(doc, "2.1 New Content - Article #18 Best Lomi LIVE", level=2, color=FOREST_GREEN)
    add_body(doc, "Article #18 published Mon Jun 29, 2026, on schedule. Verified live from the post-sitemap and "
                  "raw HTML:")
    add_styled_table(doc,
        ["Attribute", "Verified value"],
        [
            ["HTTP status", "200"],
            ["Robots", "index, follow"],
            ["Schema", "Article + WebPage JSON-LD present"],
            ["Length", "~2,600 words (long-form guide)"],
            ["Internal links", "Property pages + Book-Direct (#6) - linking floor met"],
            ["Live slug", "best-lomi-lipa-city"],
            ["Planned slug (drift)", "best-lomi-lipa-city-local-guide (corrected in tracker)"],
        ],
        col_widths=[2.2, 4.8])
    add_body(doc, "Slug drift is the recurring Yoast-shortening pattern (same as #17). The tracker Target URL and "
                  "Confirmed Live Slug were synced to the live value; status set to 'Live - verify GSC'.", italic=True)

    add_heading_styled(doc, "2.2 Publishing Cadence - On Schedule", level=2, color=FOREST_GREEN)
    add_styled_table(doc,
        ["#", "Article", "Target Monday", "Status"],
        [
            ["17", "Lipa Barako Coffee Heritage", "Jun 22", "LIVE (indexed)"],
            ["18", "Best Lomi in Lipa", "Jun 29", "LIVE (verify GSC)"],
            ["19", "Work From Lipa (Rainy Season)", "Jul 6", "Not due yet - on schedule"],
        ],
        col_widths=[0.4, 2.8, 1.2, 2.6])
    add_body(doc, "#19's planned URL correctly returns 404/noindex at audit time - it is scheduled for Jul 6 and "
                  "is the very next publish. No action.", italic=True)

    add_heading_styled(doc, "2.3 Main-Site Redeploy (Jul 2) - Schema Re-Verified", level=2, color=FOREST_GREEN)
    add_body(doc, "The main site was redeployed Jul 2, 2026 (homepage, /faq, /about, /privacy, /terms lastmod "
                  "re-stamped in sitemap.xml). Because a redeploy is a common silent-breakage point for structured "
                  "data, all schema was re-verified against raw HTML:")
    add_styled_table(doc,
        ["Page", "Schema re-verified Jul 4"],
        [
            ["/faq", "FAQPage + LocalBusiness - present"],
            ["/properties/cozy-1-bedroom", "VacationRental + Accommodation + FAQPage - present"],
            ["/properties/spacious-2-bedroom", "VacationRental + Accommodation + FAQPage - present"],
            ["3x Mickey property pages (7/11/15)", "VacationRental + Accommodation + FAQPage - present"],
        ],
        col_widths=[3.4, 3.6])
    add_body(doc, "Occupancy convention also re-confirmed on-page (sleeps-7: P2,400, 'covers 5 / sleeps up to 7'). "
                  "No regressions from the redeploy.", italic=True)

    add_heading_styled(doc, "2.4 Archive noindex - Now Confirmed Live", level=2, color=FOREST_GREEN)
    add_body(doc, "The tag/author archive noindex directive (set ~Jun 2; the Jul 3 note flagged the archives were "
                  "still indexed because they had not yet been recrawled) is now confirmed present in raw HTML on "
                  "all four archives. They will drop from the index on Google's next crawl:")
    add_bullet(doc, "/tag/airbnb/ -> noindex, follow")
    add_bullet(doc, "/tag/batangas/ -> noindex, follow")
    add_bullet(doc, "/tag/barako-coffee/ -> noindex, follow")
    add_bullet(doc, "/author/cassandrakim/ -> noindex, follow")
    add_body(doc, "Tracker status for these rows was reconciled to 'noindexed' (three previously read 'Indexed').",
             italic=True)

    add_heading_styled(doc, "2.5 NEW & RESOLVED - /book Endpoints Canonicalized to Homepage", level=2, color=FOREST_GREEN)
    add_body(doc, "New finding this run: the three Mickey booking-form endpoints returned 200 but served the "
                  "HOMEPAGE's canonical, title, and og:url (SPA app-shell with no per-route rendering). They were "
                  "not disallowed in robots.txt, and GSC had indexed two of them (sleeps-7/book, sleeps-15/book; "
                  "sleeps-11/book was not indexed). Effect: thin booking-form routes all claiming to be the "
                  "homepage - diluting homepage signal consolidation and cluttering the index with near-duplicates.")
    add_body(doc, "RESOLVED same-day. The recommended fix was implemented and verified live on all three routes:",
             bold=True)
    add_styled_table(doc,
        ["Route", "robots", "canonical", "title"],
        [
            ["sleeps-7/book", "noindex, follow", "self", "unique"],
            ["sleeps-11/book", "noindex, follow", "self", "unique"],
            ["sleeps-15/book", "noindex, follow", "self", "unique"],
        ],
        col_widths=[2.4, 1.6, 1.2, 1.8])
    add_body(doc, "noindex (not robots.txt Disallow) was used deliberately, so Google can still crawl the pages to "
                  "SEE the directive and drop the two already-indexed ones. GSC recrawl was requested to force the "
                  "drop; monitor Coverage over the next 1-2 weeks. This also resolves the standing 'decide on /book "
                  "endpoints' open item.", italic=True)
    add_divider(doc)

    # ---------------- 3. TRACKER RECONCILIATION ----------------
    add_heading_styled(doc, "3. Keyword Tracker Reconciliation", level=1, color=CORAL)
    add_body(doc, "Workbook backed up before edits as HavenInLipa_SEO_Tracker.pre-070426delta.backup.xlsx. "
                  "All 12 sheets and the formula-driven Cluster KPIs tab verified intact after the edits.")

    add_heading_styled(doc, "3.1 Article #18 - Slug Sync & Status Flip", level=2, color=FOREST_GREEN)
    add_bullet(doc, "Page Indexing Tracker: #18 row URL updated to the live slug (best-lomi-lipa-city); status "
                    "flipped 'Page Not Published' -> 'Live - verify GSC'; action set to 'Request indexing in GSC'.")
    add_bullet(doc, "Keyword Master: all three #18 rows (primary 'best lomi in lipa' + two secondaries) had Target "
                    "URL and Confirmed Live Slug synced to the live value.")

    add_heading_styled(doc, "3.2 Archive Rows Reconciled", level=2, color=FOREST_GREEN)
    add_body(doc, "The four archive rows (/tag/barako-coffee, /tag/batangas, /tag/airbnb, /author/cassandrakim) "
                  "were set to a consistent 'noindexed' status with a 'confirmed live in raw HTML 2026-07-04 - "
                  "pending drop' note. Previously three read 'Indexed' and only one read 'noindexed'.")

    add_heading_styled(doc, "3.3 /book Rows Updated", level=2, color=FOREST_GREEN)
    add_body(doc, "The two indexed /book rows (sleeps-7/book, sleeps-15/book) were updated to 'noindex - pending "
                  "drop', with the fix recorded (noindex + self-canonical + unique title) and the action set to "
                  "'GSC recrawl requested; monitor Coverage for drop'.")

    add_heading_styled(doc, "3.4 Continuity Note", level=2, color=FOREST_GREEN)
    add_body(doc, "No net-new keywords or strategy were introduced. All edits are hygiene/reconciliation against "
                  "the existing 6-cluster tracker - the single source of truth per the standing continuity rule.")
    add_divider(doc)

    # ---------------- 4. COMPETITOR POSITION ----------------
    add_heading_styled(doc, "4. Competitor Position (Carry-Forward)", level=1, color=CORAL)
    add_body(doc, "Competitor re-analysis was intentionally out of scope for this delta (moat unchanged, baseline "
                  "< 90 days). The prior finding stands: no local accommodation competitor (Lakeview Resort, JET "
                  "Hotel, The Farm at San Benito, Cintai Corito's Garden) runs a content/SEO engine on planning or "
                  "'things to do' intent. HavenInLipa remains the only Lipa lodging brand ranking on informational "
                  "and planning intent, widened further by the branded Mickey-themed-rental terms that have no "
                  "local competitor at all. Available on request if a refresh is wanted.")
    add_divider(doc)

    # ---------------- 5. FINDINGS & RESOLUTION LOG ----------------
    add_heading_styled(doc, "5. Findings & Resolution Log", level=1, color=CORAL)
    add_styled_table(doc,
        ["Finding (Jul 4)", "Severity", "Status", "Resolution / Owner"],
        [
            ["#18 Best Lomi live but tracker said 'Not Published' + slug drift", "Med (hygiene)", "RESOLVED",
             "Tracker URL/slug synced; status flipped to Live. NCS."],
            ["/book endpoints canonicalize to homepage; 2 indexed", "Med", "FIXED",
             "noindex,follow + self-canonical + unique titles shipped & verified; GSC recrawl requested. Client dev."],
            ["Archive noindex not yet confirmed effective", "Low", "CONFIRMED LIVE",
             "noindex,follow verified in raw HTML on all 4 archives; pending drop on next crawl."],
            ["Main-site redeploy (Jul 2) - schema regression risk", "-", "HOLDING",
             "All property + FAQ JSON-LD re-verified present. No regressions."],
            ["#19 planned URL returns 404", "-", "EXPECTED",
             "Scheduled Jul 6; not yet due. No action."],
            ["GSC indexing of #18 + newly-live pages", "Low", "OPEN",
             "Request indexing / confirm in next GSC pull. Client."],
        ],
        col_widths=[2.6, 0.9, 1.0, 2.5])
    add_divider(doc)

    # ---------------- 6. RECOMMENDATIONS ----------------
    add_heading_styled(doc, "6. Recommendations & Next Steps", level=1, color=CORAL)

    add_heading_styled(doc, "6.1 Immediate (Client-Owned)", level=2, color=FOREST_GREEN)
    add_bullet(doc, "Confirm the two indexed /book endpoints drop from GSC Coverage over the next 1-2 weeks now "
                    "that noindex + recrawl request are in place.")
    add_bullet(doc, "Request indexing / confirm GSC indexing for #18 Best Lomi.")
    add_bullet(doc, "Confirm the four tag/author archives fall out of the index once Google recrawls them.")
    add_bullet(doc, "Still open from prior runs: supply per-extra-guest fees for each Mickey config (Melody) to "
                    "document 'sleeps up to N (rate for M)' on the property pages.")

    add_heading_styled(doc, "6.2 Continue the Plan", level=2, color=FOREST_GREEN)
    add_bullet(doc, "Publish #19 Work From Lipa (Rainy Season) on Mon Jul 6, then continue weekly through #26 "
                    "(Aug 24); re-confirm each live slug against the sitemap at publish (recurring drift point).")
    add_bullet(doc, "Run the R1 'Things to Do' quarterly refresh as pre-scheduled (Aug 31).")
    add_bullet(doc, "Publish gap articles #27/#28/#29 in September; after #28, wire the hub links from "
                    "#1/#4/#5/#7/#10.")
    add_bullet(doc, "Continue baselining the KPI Dashboard - fill the remaining GSC/GBP rows for May-Jul.")

    add_heading_styled(doc, "6.3 KPIs to Track", level=2, color=FOREST_GREEN)
    add_styled_table(doc,
        ["KPI", "Source", "Cadence"],
        [
            ["/book + archive drop-out from index", "GSC Coverage / URL Inspection", "Weekly until dropped"],
            ["Indexing status of #18 + newly-live pages", "GSC URL Inspection", "Per publish / weekly"],
            ["Impressions / clicks per tracked keyword", "GSC Performance", "Weekly (Mon)"],
            ["Rich-result eligibility (VacationRental/FAQ)", "GSC Enhancements", "Monthly"],
            ["Indexed page count trend (20->25->31)", "GSC Indexing / KPI Dashboard", "Monthly"],
        ],
        col_widths=[3.2, 2.3, 1.5])
    add_divider(doc)

    closing = doc.add_paragraph()
    closing.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = closing.add_run("Prepared by NetCoreSolutions  |  July 4, 2026  |  Delta Audit (070426)  |  CONFIDENTIAL")
    run.font.size = Pt(9)
    run.font.color.rgb = MED_GRAY
    run.font.name = "Calibri"

    doc.save(OUTPUT_FILE)
    doc.save(OUTPUT_FILE_LATEST)
    os.makedirs(os.path.dirname(OUTPUT_FILE_RUN), exist_ok=True)
    doc.save(OUTPUT_FILE_RUN)
    print("Saved:", OUTPUT_FILE)
    print("Saved (latest):", OUTPUT_FILE_LATEST)
    print("Saved (run copy):", OUTPUT_FILE_RUN)


if __name__ == "__main__":
    generate_report()
