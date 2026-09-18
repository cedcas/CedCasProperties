#!/usr/bin/env python3
"""
Apply Option 2 to the HavenInLipa Keyword Tracker:
1. Remove the 'Indexing Status' column (col G) from Keyword Master.
2. Correct blog URLs in column E to use https://blog.haveninlipa.com/ subdomain.
3. Add a new 'Page Indexing Tracker' sheet pre-populated with all unique URLs
   and indexing status from the May 2026 GSC screenshot.
Preserves all user-entered values in other tracking columns.
"""

import os
import re
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

BASE_DIR = "/Volumes/Files and Cloud/Dropbox/VSCode (non-Git)/seo"
FILE = os.path.join(BASE_DIR, "content for HavenInLipa", "HavenInLipa_Keyword_Tracker.xlsx")

CORAL = "FF5371"
FOREST_GREEN = "3B5323"
LIGHT_GREEN = "E8EFE0"
LIGHT_CORAL = "FFE4E9"
GRAY = "F2F2F2"
DARK_TEXT = "2C2C2C"
WHITE = "FFFFFF"
INDEXED_GREEN = "0D8A4E"
PENDING_ORANGE = "E88A1A"
NOT_PUBLISHED_RED = "D93B3B"

THIN = Side(border_style="thin", color="CCCCCC")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def style_header(cell, fill=CORAL, color=WHITE):
    cell.font = Font(name="Calibri", bold=True, color=color, size=11)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = BORDER


def style_body(cell, bold=False, fill=None, color=DARK_TEXT, align="left"):
    cell.font = Font(name="Calibri", bold=bold, color=color, size=10)
    if fill:
        cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=True)
    cell.border = BORDER


def status_pill(cell, status):
    s = status.strip().lower()
    if s == "indexed":
        fill, color = INDEXED_GREEN, WHITE
    elif "requested" in s or "discovered" in s:
        fill, color = PENDING_ORANGE, WHITE
    elif "not published" in s or "crawled not indexed" in s:
        fill, color = NOT_PUBLISHED_RED, WHITE
    elif "n/a" in s:
        fill, color = "888888", WHITE
    else:
        return
    cell.font = Font(name="Calibri", bold=True, color=color, size=10)
    cell.fill = PatternFill("solid", fgColor=fill)
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


# ---------------------------------------------------------------------------
# 1. Remove Indexing Status column (col G) from Keyword Master
# ---------------------------------------------------------------------------

def remove_indexing_status_column(wb):
    ws = wb["Keyword Master"]
    header = ws.cell(row=1, column=7).value
    if header == "Indexing Status":
        ws.delete_cols(7)
        print("Removed 'Indexing Status' column from Keyword Master")
    else:
        print(f"WARNING: column G header is {header!r}, not 'Indexing Status'. Skipping delete.")


# ---------------------------------------------------------------------------
# 2. Update Target URL / Page values to use real subdomain structure
# ---------------------------------------------------------------------------

# Hard-mapped special cases (slug differs from what we originally guessed)
URL_HARDMAP = {
    "/ (Homepage)": "https://haveninlipa.com/",
    "/ + Property pages": "https://haveninlipa.com/ + property pages",
    "/blog/holy-week-getaway-lipa": "https://blog.haveninlipa.com/holy-week-in-lipa-city-batangas-a-peaceful-retreat-near-manila/",
    "/properties/mickey-in-lipa": "https://haveninlipa.com/properties/mickey-in-lipa",
}

def normalize_url(value):
    if not value:
        return value
    v = str(value).strip()
    if v in URL_HARDMAP:
        return URL_HARDMAP[v]
    # /blog/<slug>  -> https://blog.haveninlipa.com/<slug>/
    m = re.match(r"^/blog/([a-z0-9\-]+)/?$", v)
    if m:
        return f"https://blog.haveninlipa.com/{m.group(1)}/"
    # Bare /properties/<slug>
    m = re.match(r"^/properties/([a-z0-9\-]+)/?$", v)
    if m:
        return f"https://haveninlipa.com/properties/{m.group(1)}"
    # Already a full URL, or TBD/gap placeholder, leave alone
    return v


def fix_target_urls(wb):
    ws = wb["Keyword Master"]
    changed = 0
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(row=r, column=5)
        old = cell.value
        new = normalize_url(old)
        if new != old:
            cell.value = new
            changed += 1
    print(f"Updated {changed} Target URL values in Keyword Master")


# ---------------------------------------------------------------------------
# 3. Build Page Indexing Tracker sheet
# ---------------------------------------------------------------------------

# (URL, Page Type, Linked Articles / Keywords, Indexing Status, Last Crawled, Action Needed, Notes)
PAGES = [
    # Indexed (confirmed via GSC May 2026 screenshot)
    ("https://haveninlipa.com/",
     "Homepage", "Brand SERP + Cluster 1 transactional", "Indexed", "2026-05-11",
     "Monitor weekly for brand term ranking",
     "Verified indexed (GSC May 11, 2026)"),
    ("https://www.haveninlipa.com/",
     "Homepage (www)", "Same as apex", "Indexed", "2026-04-18",
     "Confirm canonical points to https://haveninlipa.com/",
     "Both apex and www are indexed - check for duplicate content issue"),
    ("https://haveninlipa.com/properties/spacious-2-bedroom",
     "Property page", "Cluster 1 (transactional)", "Indexed", "2026-05-08",
     "Expand to 1,500+ words per May 8 strategy report",
     "Verified indexed (GSC May 8, 2026)"),
    ("https://haveninlipa.com/properties/cozy-1-bedroom",
     "Property page", "Cluster 1 (transactional) + couple keywords", "Indexed", "2026-05-08",
     "Expand to 1,500+ words; feature 1BR for couples",
     "Verified indexed (GSC May 8, 2026)"),
    ("https://haveninlipa.com/about",
     "About page", "Brand SERP", "Indexed", "2026-05-08",
     "Add author credentials and local Lipa story for E-E-A-T",
     "Verified indexed (GSC May 8, 2026)"),
    ("https://haveninlipa.com/faq",
     "FAQ page", "Brand SERP + multiple long-tail", "Indexed", "2026-05-08",
     "Add FAQ schema markup for rich results",
     "Verified indexed (GSC May 8, 2026)"),
    ("https://haveninlipa.com/privacy",
     "Legal", "N/A", "Indexed", "2026-05-08",
     "No action needed",
     "Verified indexed (GSC May 8, 2026)"),
    ("https://haveninlipa.com/terms",
     "Legal", "N/A", "Indexed", "2026-05-08",
     "No action needed",
     "Verified indexed (GSC May 8, 2026)"),
    ("https://blog.haveninlipa.com/",
     "Blog index", "All blog clusters", "Indexed", "2026-04-11",
     "Ensure latest article is featured at top; sticky 'Book Direct' CTA",
     "Verified indexed (GSC Apr 11, 2026)"),
    ("https://blog.haveninlipa.com/category/booking-tips/",
     "Blog category", "Cluster 5 (direct booking)", "Indexed", "2026-04-30",
     "Add intro copy to category page (currently thin)",
     "Verified indexed (GSC Apr 30, 2026)"),
    ("https://blog.haveninlipa.com/holy-week-in-lipa-city-batangas-a-peaceful-retreat-near-manila/",
     "Blog article #8", "holy week getaway lipa", "Indexed", "2026-05-06",
     "Reframe to evergreen + queue for Holy Week 2027 push",
     "Verified indexed (GSC May 6, 2026)"),

    # Requested for indexing (per user: ~50+ URLs requested in past few days)
    ("https://blog.haveninlipa.com/things-to-do-lipa-city/",
     "Blog article #1", "Cluster 2 (things to do)", "Not Indexed - Requested", "",
     "Wait 1-7 days, then URL Inspect to confirm",
     "Targets: things to do in lipa city, lipa city tourist spots"),
    ("https://blog.haveninlipa.com/weekend-getaway-lipa/",
     "Blog article #3", "Cluster 3 (weekend getaway)", "Not Indexed - Requested", "",
     "Wait 1-7 days, then URL Inspect to confirm",
     "Targets: weekend getaway near manila"),
    ("https://blog.haveninlipa.com/mt-maculot-hiking-guide/",
     "Blog article #4", "Cluster 2 (informational)", "Not Indexed - Requested", "",
     "Wait 1-7 days, then URL Inspect to confirm",
     "Targets: mt maculot hiking guide"),
    ("https://blog.haveninlipa.com/best-restaurants-lipa/",
     "Blog article #5", "Cluster 4 (food/travel planning)", "Not Indexed - Requested", "",
     "Wait 1-7 days, then URL Inspect to confirm",
     "Targets: best restaurants in lipa, lipa city food trip"),
    ("https://blog.haveninlipa.com/why-book-direct-vs-airbnb/",
     "Blog article #6", "Cluster 5 (direct booking) - HIGH priority", "Not Indexed - Requested", "",
     "Wait 1-7 days, then URL Inspect to confirm",
     "Targets: book direct vs airbnb, save on vacation rental fees, airbnb alternative lipa"),
    ("https://blog.haveninlipa.com/taal-volcano-day-trip/",
     "Blog article #7", "Cluster 2 (day trip)", "Not Indexed - Requested", "",
     "Wait 1-7 days, then URL Inspect to confirm",
     "Targets: taal volcano day trip from lipa"),
    ("https://blog.haveninlipa.com/remote-work-staycation-lipa/",
     "Blog article #9", "Cluster 3 (remote work)", "Not Indexed - Requested", "",
     "Wait 1-7 days, then URL Inspect to confirm",
     "Targets: remote work staycation philippines"),
    ("https://blog.haveninlipa.com/how-to-get-to-lipa-from-manila/",
     "Blog article #10", "Cluster 4 (logistics)", "Not Indexed - Requested", "",
     "Confirm article is actually live (was reported 404); then verify indexing",
     "Targets: how to get to lipa from manila, lipa city travel guide"),
    ("https://blog.haveninlipa.com/family-staycation-lipa/",
     "Blog article #11", "Cluster 3 (family) - HIGH priority", "Not Indexed - Requested", "",
     "Wait 1-7 days, then URL Inspect to confirm",
     "Targets: family staycation lipa, family-friendly rental lipa city, weekend getaway with kids near manila"),
    ("https://blog.haveninlipa.com/romantic-getaway-batangas/",
     "Blog article #12", "Cluster 3 (couples) - HIGH priority", "Not Indexed - Requested", "",
     "Wait 1-7 days, then URL Inspect to confirm",
     "Targets: romantic getaway batangas, couple retreat lipa, honeymoon batangas, anniversary getaway near manila"),
    ("https://blog.haveninlipa.com/coming-soon-disney-inspired-family-house-lipa/",
     "Blog article #13", "Cluster 6 (Mickey in Lipa launch)", "Not Indexed - Requested", "",
     "Wait 1-7 days, then URL Inspect to confirm",
     "Targets: mickey in lipa, disney-inspired house lipa, full house rental lipa"),
    ("https://blog.haveninlipa.com/family-weekend-batangas-without-beach-crowds/",
     "Blog article #14", "Cluster 3 (family) - NEW", "Not Indexed - Requested", "",
     "Wait 1-7 days, then URL Inspect to confirm",
     "Targets: family weekend batangas, weekend near manila with family, batangas without beach crowds, lipa family itinerary"),
    ("https://haveninlipa.com/properties/mickey-in-lipa",
     "Property page (NEW)", "Mickey in Lipa launch", "Not Indexed - Requested", "",
     "Confirm property page is built; this is the booking URL for Article #13",
     "Targets: mickey in lipa, disney-inspired house lipa, bella vita lipa rental"),

    # Pages not yet built (Tier C / D gaps from May 8 strategy)
    ("(TBD) https://blog.haveninlipa.com/barkada-group-trip-lipa/",
     "Future article (Tier C)", "Cluster 3 (barkada)", "Page Not Published", "",
     "Write and publish during Tier C window (June 2026)",
     "Targets: barkada group trip lipa"),
    ("(TBD) https://blog.haveninlipa.com/summer-in-lipa-2026/",
     "Future seasonal LP", "Cluster 3 (seasonal)", "Page Not Published", "",
     "Publish before May-June 2026 peak; landing page format",
     "Targets: summer in lipa 2026"),
    ("(TBD) https://blog.haveninlipa.com/batangas-road-trip-itinerary/",
     "Future hub article", "Cluster 4 (travel planning)", "Page Not Published", "",
     "Write during Tier C window; will interlink with #1, #4, #5, #7, #10",
     "Targets: batangas road trip itinerary"),
    ("(TBD) https://blog.haveninlipa.com/senior-friendly-lipa-getaway/",
     "Future article (Tier D)", "Cluster 3 (audience)", "Page Not Published", "",
     "Write during Q3 2026 authority phase",
     "Targets: senior-friendly lipa getaway"),
    ("(TBD) https://blog.haveninlipa.com/solo-travel-guide-lipa/",
     "Future article (Tier D)", "Cluster 3 (audience)", "Page Not Published", "",
     "Write during Q3 2026 authority phase",
     "Targets: solo travel guide to lipa"),
]


def build_page_indexing_tracker(wb):
    # Remove existing sheet if it exists (allow re-run)
    if "Page Indexing Tracker" in wb.sheetnames:
        del wb["Page Indexing Tracker"]
    ws = wb.create_sheet("Page Indexing Tracker")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "B2"

    headers = [
        "URL",
        "Page Type",
        "Linked Cluster / Keywords",
        "Indexing Status",
        "Last Crawled (GSC)",
        "Last Checked (you)",
        "Action Needed",
        "Notes",
    ]
    for i, h in enumerate(headers, start=1):
        style_header(ws.cell(row=1, column=i, value=h), fill=FOREST_GREEN)
    ws.row_dimensions[1].height = 38

    for r, page in enumerate(PAGES, start=2):
        url, ptype, linked, status, last_crawled, action, notes = page
        ws.cell(row=r, column=1, value=url)
        ws.cell(row=r, column=2, value=ptype)
        ws.cell(row=r, column=3, value=linked)
        ws.cell(row=r, column=4, value=status)
        ws.cell(row=r, column=5, value=last_crawled)
        ws.cell(row=r, column=6, value="")  # Last Checked (user fills weekly)
        ws.cell(row=r, column=7, value=action)
        ws.cell(row=r, column=8, value=notes)

        zebra = LIGHT_GREEN if r % 2 == 0 else WHITE
        for col in range(1, len(headers) + 1):
            style_body(ws.cell(row=r, column=col), fill=zebra)

        # Status pill on column 4
        status_pill(ws.cell(row=r, column=4), status)
        ws.row_dimensions[r].height = 34

    widths = [56, 22, 38, 24, 18, 18, 42, 50]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # Reorder so this sheet sits right after Priority Focus
    target_order = [
        "README",
        "Keyword Master",
        "Priority Focus (60-Day)",
        "Page Indexing Tracker",
        "GSC Setup Steps",
        "GSC Troubleshooting",
        "Blogger Brief",
        "Tracking Log Template",
    ]
    existing = wb.sheetnames
    new_order = [name for name in target_order if name in existing]
    for name in existing:
        if name not in new_order:
            new_order.append(name)
    wb._sheets = [wb[name] for name in new_order]


# ---------------------------------------------------------------------------
# MAIN
# ---------------------------------------------------------------------------

def main():
    if not os.path.exists(FILE):
        raise SystemExit(f"File not found: {FILE}")

    wb = load_workbook(FILE)
    print(f"Loaded sheets: {wb.sheetnames}")

    remove_indexing_status_column(wb)
    fix_target_urls(wb)
    build_page_indexing_tracker(wb)

    wb.save(FILE)
    print(f"Saved: {FILE}")
    print(f"Final sheets: {wb.sheetnames}")


if __name__ == "__main__":
    main()
