"""Recolor Page Indexing Tracker status cells whose text changed but whose
fill color is stale from the previous status (openpyxl doesn't do this
automatically — cell fills are static, not conditional formatting)."""
import copy
import openpyxl

WORKBOOK = "../HavenInLipa_SEO_Tracker.xlsx"

GREEN_URLS = {
    "https://blog.haveninlipa.com/barkada-getaway-near-manila-why-a-whole-house-in-lipa-beats-a-beach-resort-cost-per-head-math-inside/",
    "https://blog.haveninlipa.com/indoor-things-to-do-in-lipa-city-when-it-rains/",
    "https://blog.haveninlipa.com/best-lomi-lipa-city/",
    "https://blog.haveninlipa.com/work-from-lipa-rainy-season-july/",
    "https://blog.haveninlipa.com/lipa-vs-tagaytay-an-honest-comparison-from-a-host-who-lives-in-lipa/",
    "https://blog.haveninlipa.com/casa-de-segunda-lipa-city/",
    "https://blog.haveninlipa.com/august-long-weekends-lipa-city-2026/",
    "https://blog.haveninlipa.com/lipa-charter-day-history-coffee-town-to-city/",
    "https://blog.haveninlipa.com/ninoy-aquino-day-long-weekend-lipa-2026/",
    "https://blog.haveninlipa.com/tag/batangas/",
    "https://blog.haveninlipa.com/tag/airbnb/",
    "https://blog.haveninlipa.com/author/cassandrakim/",
}

RED_URLS = {
    "https://blog.haveninlipa.com/category/booking-tips/",
    "https://blog.haveninlipa.com/tag/barako-coffee/",
    "https://blog.haveninlipa.com/category/travel-and-itineraries/",
    "https://blog.haveninlipa.com/category/weekend-getaways/",
    "https://blog.haveninlipa.com/category/outdoor-adventures/",
    "https://blog.haveninlipa.com/category/uncategorized/",
    "https://blog.haveninlipa.com/category/getting-here/",
}


def main():
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb["Page Indexing Tracker"]
    headers = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(headers)}
    status_col = col["Indexing Status"]

    green_exemplar = ws.cell(2, status_col)   # https://haveninlipa.com/ -> Indexed
    red_exemplar = None
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, col["URL"]).value == "https://blog.haveninlipa.com/lipa-pilgrimage-guide/":
            red_exemplar = ws.cell(r, status_col)
            break
    assert red_exemplar is not None

    recolored = []
    for r in range(2, ws.max_row + 1):
        url = ws.cell(r, col["URL"]).value
        cell = ws.cell(r, status_col)
        if url in GREEN_URLS:
            cell.fill = copy.copy(green_exemplar.fill)
            cell.font = copy.copy(green_exemplar.font)
            recolored.append((url, "green"))
        elif url in RED_URLS:
            cell.fill = copy.copy(red_exemplar.fill)
            cell.font = copy.copy(red_exemplar.font)
            recolored.append((url, "red"))

    wb.save(WORKBOOK)
    print(f"Recolored {len(recolored)} status cells")
    for url, color in recolored:
        print(f"  {color}: {url}")


if __name__ == "__main__":
    main()
